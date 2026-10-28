"""学校规范与 DOCX 属性检查的离线预埋对照。"""
import io
import json
import math
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.enum.section import WD_SECTION
from flask import Flask

from school_format import school_format_bp, _style_chain


def docx_fixture(*, size=12, inherited=True, section_margin=2.54, latin="Times New Roman"):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.font.name = latin
    rpr = normal._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(0.74)
    doc.sections[0].top_margin = Cm(section_margin)
    doc.add_paragraph("封面标题")
    doc.add_paragraph("第一章 绪论", style="Heading 1")
    para = doc.add_paragraph()
    zh = para.add_run("这是一段中文正文")
    en = para.add_run(" English body")
    if not inherited:
        zh.font.size = Pt(size)
        en.font.size = Pt(size)
    doc.add_paragraph("参考文献", style="Heading 1")
    doc.add_paragraph("中文参考文献不应受正文规则限制")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def rule(field, expected, source="school", sha=None, quote=None, **kwargs):
    item = {"field": field, "expected": expected, "source_type": source,
            "confirmed": True, "active": True}
    if sha:
        item["source_hash"] = sha
    if quote:
        item["quote"] = quote
    item.update(kwargs)
    return item


class SchoolFormatTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app = Flask(__name__)
        app.register_blueprint(school_format_bp)
        cls.app = app

    def setUp(self):
        self.client = self.app.test_client()
        self.requirement = b"Body font: Song; 12 pt. Margins: 2.54 cm."
        preview = self.client.post("/api/school_rules/preview", data={
            "requirements": (io.BytesIO(self.requirement), "school.txt")})
        self.assertEqual(preview.status_code, 200)
        self.sha = preview.json["sources"][0]["sha256"]

    def profile(self, rules):
        return {"confirmed": True, "scope": {"school": "Test University", "year": "2026",
                                               "degree": "本科", "body": True}, "rules": rules}

    def check(self, paper, profile, requirements=True, filename="paper.docx"):
        data = {"paper": (io.BytesIO(paper), filename),
                "profile": json.dumps(profile, ensure_ascii=False)}
        if requirements:
            data["requirements"] = (io.BytesIO(self.requirement), "school.txt")
        return self.client.post("/api/format_check", data=data)

    def test_inherited_mixed_font_size_line_indent_and_margin(self):
        rules = [rule("east_asia_font", "宋体", sha=self.sha, quote="Body font: Song"),
                 rule("latin_font", "Times New Roman", source="custom"),
                 rule("font_size_pt", 12, sha=self.sha, quote="12 pt"),
                 rule("line_spacing_multiplier", 1.5, source="custom"),
                 rule("first_line_indent_cm", 0.74, source="custom"),
                 rule("margin_top_cm", 2.54, sha=self.sha, quote="2.54 cm")]
        response = self.check(docx_fixture(), self.profile(rules))
        self.assertEqual(response.status_code, 200, response.json)
        report = response.json
        self.assertEqual(report["counts"]["noncompliant"], 0, report["findings"])
        self.assertEqual(report["counts"]["compliant"], 7, report["findings"])
        self.assertTrue(all("第3段" in x["location"] or "第1节" in x["location"]
                            for x in report["findings"]))
        self.assertEqual(report["coverage"]["body_paragraphs_checked"], 1)
        self.assertIn("论文 SHA256", report["markdown"])

    def test_independent_mutations_are_located(self):
        rules = [rule("font_size_pt", 12, sha=self.sha, quote="12 pt"),
                 rule("margin_top_cm", 2.54, sha=self.sha, quote="2.54 cm")]
        result = self.check(docx_fixture(size=14, inherited=False, section_margin=3.0),
                            self.profile(rules)).json
        self.assertEqual(result["counts"]["noncompliant"], 3, result["findings"])
        self.assertTrue(any(f["actual"] == 14 and "第3段" in f["location"] for f in result["findings"]))
        self.assertTrue(any(f["field"] == "margin_top_cm" and abs(f["actual"] - 3.0) < 0.005
                            for f in result["findings"]))

    def test_school_conflict_and_custom_override_cannot_pass(self):
        a = rule("font_size_pt", 12, sha=self.sha, quote="12 pt")
        b = rule("font_size_pt", 14, sha=self.sha, quote="Body font: Song")
        report = self.check(docx_fixture(), self.profile([a, b])).json
        self.assertEqual(report["counts"]["conflict"], 1)
        self.assertEqual(report["counts"]["compliant"], 0)
        override = rule("font_size_pt", 14, source="custom", reason="导师要求")
        report = self.check(docx_fixture(), self.profile([a, override])).json
        self.assertEqual(report["counts"]["conflict"], 1)
        self.assertEqual(report["counts"]["compliant"], 0)
        self.assertIn("导师要求", report["findings"][0]["reason"])

    def test_provenance_is_rechecked_per_request(self):
        school = rule("font_size_pt", 12, sha=self.sha, quote="12 pt")
        missing = self.check(docx_fixture(), self.profile([school]), requirements=False)
        self.assertEqual(missing.status_code, 400)
        fake = rule("font_size_pt", 12, sha=self.sha, quote="16 pt")
        self.assertEqual(self.check(docx_fixture(), self.profile([fake])).status_code, 400)
        unconfirmed = self.profile([school])
        unconfirmed["confirmed"] = False
        self.assertEqual(self.check(docx_fixture(), unconfirmed).status_code, 400)

    def test_pdf_and_unknown_docx_properties_are_undetermined(self):
        from pypdf import PdfWriter
        stream = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        writer.write(stream)
        profile = self.profile([rule("font_size_pt", 12, source="custom")])
        pdf = self.check(stream.getvalue(), profile, requirements=False, filename="paper.pdf")
        self.assertEqual(pdf.json["counts"]["undetermined"], 1)
        doc = Document()
        doc.add_paragraph("第一章 绪论", style="Heading 1")
        doc.add_paragraph("这是一段中文正文")
        # ASCII 字体在默认主题上不能作为学校指定字体的确定证据。
        out = io.BytesIO()
        doc.save(out)
        unknown = self.check(out.getvalue(), self.profile([
            rule("east_asia_font", "宋体", source="custom")]), requirements=False)
        self.assertEqual(unknown.json["counts"]["undetermined"], 1, unknown.json)

    def test_distinct_sections_and_closed_rule(self):
        doc = Document(io.BytesIO(docx_fixture()))
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        section.top_margin = Cm(3.0)
        output = io.BytesIO()
        doc.save(output)
        rules = [rule("margin_top_cm", 2.54, sha=self.sha, quote="2.54 cm"),
                 rule("font_size_pt", 12, source="custom", active=False)]
        result = self.check(output.getvalue(), self.profile(rules)).json
        self.assertEqual(result["counts"]["compliant"], 1)
        self.assertEqual(result["counts"]["noncompliant"], 1)
        self.assertEqual(result["counts"]["not_applicable"], 1)
        self.assertIn("第2节", result["findings"][1]["location"])

    def test_bad_and_oversized_files_rejected(self):
        self.assertEqual(self.check(b"bad", self.profile([rule("font_size_pt", 12, source="custom")]),
                                    requirements=False).status_code, 400)
        response = self.client.post("/api/school_rules/preview", data={
            "requirements": (io.BytesIO(b"x" * (8 * 1024 * 1024 + 1)), "too-big.txt")})
        self.assertEqual(response.status_code, 400)
        invalid_pdf = self.check(b"not-pdf", self.profile([rule("font_size_pt", 12, source="custom")]),
                                 requirements=False, filename="paper.pdf")
        self.assertEqual(invalid_pdf.status_code, 400)

    def test_cyclic_style_does_not_hang_and_character_indent_abstains(self):
        doc = Document()
        doc.add_paragraph("第一章 绪论", style="Heading 1")
        para = doc.add_paragraph("正文段落，保留字符单位缩进。")
        normal = doc.styles["Normal"]
        # 恶意/损坏 DOCX 可能让样式 basedOn 指回自身。
        based_on = OxmlElement("w:basedOn")
        based_on.set(qn("w:val"), normal.style_id)
        normal._element.insert(0, based_on)
        self.assertLess(len(list(_style_chain(para.runs[0], para, doc))), 5)
        para.paragraph_format.first_line_indent = Cm(0.74)
        ind = para._element.pPr.find(qn("w:ind"))
        ind.set(qn("w:firstLineChars"), "200")
        buf = io.BytesIO()
        doc.save(buf)
        profile = self.profile([rule("first_line_indent_cm", 0.74, source="custom")])
        result = self.check(buf.getvalue(), profile, requirements=False)
        self.assertEqual(result.status_code, 200, result.json)
        self.assertEqual(result.json["counts"]["undetermined"], 1, result.json)

    def test_numbered_body_text_is_not_automatically_a_heading(self):
        doc = Document()
        doc.add_paragraph("1 样本量为三十人，不能靠这个序号认定章节开始。")
        doc.add_paragraph("后续文字还没有可靠的正文起点。")
        buf = io.BytesIO()
        doc.save(buf)
        result = self.check(buf.getvalue(), self.profile([
            rule("font_size_pt", 12, source="custom")]), requirements=False).json
        self.assertEqual(result["counts"]["undetermined"], 1)
        self.assertEqual(result["coverage"]["body_paragraphs_checked"], 0)

    def test_conflicting_templates_and_bad_flags_pause(self):
        template_a = rule("font_size_pt", 12, source="template", sha=self.sha, quote="12 pt")
        template_b = rule("font_size_pt", 14, source="template", sha=self.sha, quote="Body font: Song")
        report = self.check(docx_fixture(), self.profile([template_a, template_b])).json
        self.assertEqual(report["counts"]["conflict"], 1)
        self.assertEqual(report["counts"]["compliant"], 0)
        custom = rule("font_size_pt", 14, source="custom")
        self.assertEqual(self.check(docx_fixture(), self.profile([
            rule("font_size_pt", 12, sha=self.sha, quote="12 pt"), custom])).status_code, 400)
        for bad in ("false", None, 1):
            item = rule("font_size_pt", 12, source="custom", active=bad)
            self.assertEqual(self.check(docx_fixture(), self.profile([item]), requirements=False).status_code, 400)
        invalid_type = rule("font_size_pt", 12, source="invented")
        self.assertEqual(self.check(docx_fixture(), self.profile([invalid_type]), requirements=False).status_code, 400)
        # JSON 标准外数字 NaN 即使被解析，也不得进入有效规则。
        nan_rule = rule("font_size_pt", math.nan, source="custom")
        self.assertEqual(self.check(docx_fixture(), self.profile([nan_rule]), requirements=False).status_code, 400)


if __name__ == "__main__":
    unittest.main()
