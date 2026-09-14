"""
表格交叉核查测试（v2.11 · P3，本地计算 + 可选端到端，零 LLM API）
================================================================
覆盖总纲 P3「论文表格数字 vs 原始数据交叉核查」的行为契约：

    [A] 表格行解析（_parse_table_row）：n/M/SD、M±SD 合并格、无 n、
        无标签行、表注行
    [B] compare_table_stats：真实数字全对 / 篡改 n / 篡改 M / 篡改 SD /
        标签对不上（不硬猜）/ 双候选列（不硬猜）/ 无数值列只核 n /
        空表安全跳过 / 舍入容忍
    [C] build_audit_report 集成：报告含「表格交叉核对」小节、
        mismatch 进 comparisons 与 suggestions、条目可追问
    [D] docx 读取 + 端到端（需本地服务；未起自动 SKIP）

用法：.venv/Scripts/python.exe table_cross_test.py
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pandas as pd

from audit import _parse_table_row, compare_table_stats, build_audit_report
from extract_paper import read_paper_tables

PASS = 0
FAIL = 0
BASE = "http://127.0.0.1:5000"


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  [PASS] {name}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name} {detail}")


def make_docx(tables):
    """内存造一份 docx（含若干表格）。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph("测试论文：性别对成绩的影响。")
    for tbl in tables:
        t = doc.add_table(rows=len(tbl), cols=len(tbl[0]))
        for i, row in enumerate(tbl):
            for j, cell in enumerate(row):
                t.rows[i].cells[j].text = cell
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


class FakeStorage:
    """最小 FileStorage 替身（filename / stream / read）。"""

    def __init__(self, buf, filename):
        self.filename = filename
        self.stream = buf

    def read(self):
        return self.stream.getvalue()


# ---------------------------------------------------------------------------
print("[A] 表格行解析")
# ---------------------------------------------------------------------------
check("n | M | SD 三数行",
      _parse_table_row(["男", "15", "71.03", "3.25"])
      == {"label": "男", "n": 15, "mean": 71.03, "sd": 3.25})
check("M±SD 合并格",
      _parse_table_row(["女", "15", "87.07 ± 2.99"])
      == {"label": "女", "n": 15, "mean": 87.07, "sd": 2.99})
check("中文全角±也认", _parse_table_row(["女", "15", "87.07±2.99"]) is not None)
check("M | SD 两小数行（无 n）",
      _parse_table_row(["男", "71.03", "3.25"])
      == {"label": "男", "n": None, "mean": 71.03, "sd": 3.25})
check("无数字行返回 None", _parse_table_row(["男", "组一"]) is None)
check("无标签行返回 None", _parse_table_row(["15", "71.0", "3.2"]) is None)
check("表注行不当标签", _parse_table_row(["注：M 为均值", "15", "71.0", "3.25"]) is None)

# ---------------------------------------------------------------------------
print("\n[B] compare_table_stats（student_scores 实算锚点）")
# ---------------------------------------------------------------------------
df = pd.read_csv(os.path.join("examples", "student_scores.csv"))
GOOD = [["组别", "n", "M", "SD"],
        ["男", "15", "71.0", "3.25"],
        ["女", "15", "87.07", "2.99"]]

out = compare_table_stats([GOOD], df, value_col="score")
check("真实数字全对（6 格一致，零 mismatch）",
      out["mismatches"] == [] and out["checked"] == 6, f"实际={out}")

bad_n = [["组别", "n", "M", "SD"],
         ["男", "14", "71.0", "3.25"], ["女", "15", "87.07", "2.99"]]
out2 = compare_table_stats([bad_n], df, value_col="score")
check("篡改 n 检出（table_n）",
      len(out2["mismatches"]) == 1 and out2["mismatches"][0]["kind"] == "table_n")
check("n mismatch 条目含实算口径", "实算 n=15" in out2["mismatches"][0]["real"])

bad_m = [["组别", "n", "M", "SD"],
         ["男", "15", "75.0", "3.25"], ["女", "15", "87.07", "2.99"]]
out3 = compare_table_stats([bad_m], df, value_col="score")
check("篡改 M 检出（table_mean）",
      {m["kind"] for m in out3["mismatches"]} == {"table_mean"})

bad_sd = [["组别", "n", "M", "SD"],
          ["男", "15", "71.0", "9.99"], ["女", "15", "87.07", "2.99"]]
out4 = compare_table_stats([bad_sd], df, value_col="score")
check("篡改 SD 检出（table_sd）",
      {m["kind"] for m in out4["mismatches"]} == {"table_sd"})

out5 = compare_table_stats(
    [[["组别", "n", "M", "SD"],
      ["A", "15", "71.0", "3.25"], ["B", "15", "87.07", "2.99"]]],
    df, value_col="score")
check("标签对不上 → 不硬猜只 note",
      out5["mismatches"] == [] and out5["checked"] == 0 and len(out5["notes"]) == 1)

df2 = df.copy()
df2["性别2"] = df["gender"]  # 制造双候选
out6 = compare_table_stats([bad_n], df2, value_col="score")
check("双候选列 → 不硬猜", out6["mismatches"] == []
      and "唯一" in (out6["notes"][0] if out6["notes"] else ""))

out7 = compare_table_stats(
    [[["组别", "n", "M"], ["男", "15", "71.0"], ["女", "15", "87.07"]]],
    df, value_col=None)
check("无数值列只核 n（checked=2）", out7["mismatches"] == [] and out7["checked"] == 2)

out8 = compare_table_stats([[["1", "2", "3"], ["4", "5", "6"]]], df, value_col="score")
check("纯数字表安全跳过", out8["mismatches"] == [] and out8["notes"] == [])
out9 = compare_table_stats([], df, value_col="score")
check("空 tables → 全空", out9 == {"mismatches": [], "notes": [], "checked": 0})

tol = [["组别", "n", "M", "SD"],
       ["男", "15", "71.01", "3.25"], ["女", "15", "87.07", "2.99"]]
out10 = compare_table_stats([tol], df, value_col="score")
check("0.01 舍入差不报", out10["mismatches"] == [], f"实际={out10['mismatches']}")

# ---------------------------------------------------------------------------
print("\n[C] build_audit_report 集成")
# ---------------------------------------------------------------------------
claims = {
    "methods": [{"method_key": "independent_t", "method_label": "独立样本 T 检验",
                 "context": "采用独立样本 t 检验比较性别差异"}],
    "quantities": [], "variables": [],
}
from app import _summarize_column  # noqa: E402  (columns 契约同源)
cols = [_summarize_column(df[c]) for c in df.columns]
rep = build_audit_report(claims, df, cols, paper_tables=[GOOD, bad_n])
check("报告含「表格交叉核对」小节", "表格交叉核对" in rep["markdown"])
check("checked 计数 = 12（两表各 6 格，含 mismatch 的也计）",
      rep["table_check"]["checked"] == 12, f"实际={rep['table_check']['checked']}")
check("坏表 mismatch 进 comparisons",
      any(c.get("kind") == "table_n" for c in rep["comparisons"]))
check("mismatch 进建议（[表格核对] 前缀）",
      any("[表格核对]" in s for s in rep["suggestions"]))
check("comparison 带通用 summary（可追问）",
      all("summary" in c for c in rep["comparisons"]
          if c.get("kind", "").startswith("table_")))
rep0 = build_audit_report(claims, df, cols, paper_tables=None)
check("无表格 → 不出小节不出建议",
      "表格交叉核对" not in rep0["markdown"]
      and not any("[表格核对]" in s for s in rep0["suggestions"]))

# ---------------------------------------------------------------------------
print("\n[D] docx 读取 + 端到端（需本地服务）")
# ---------------------------------------------------------------------------
buf = make_docx([GOOD, bad_n])
tables = read_paper_tables(FakeStorage(buf, "paper.docx"))
check("docx 读出 2 张结构化表",
      len(tables) == 2 and tables[0][1] == ["男", "15", "71.0", "3.25"],
      f"实际={tables[:1]}")
check("非 docx 返回 []",
      read_paper_tables(FakeStorage(io.BytesIO(b"a,b\n1,2"), "x.txt")) == [])

import urllib.request
import json as _json


def _server_up():
    try:
        with urllib.request.urlopen(BASE + "/health", timeout=3) as r:
            return r.status == 200
    except Exception:  # noqa: BLE001
        return False


if _server_up():
    boundary = "----tblcross"
    paper_bytes = make_docx([GOOD, bad_n]).getvalue()
    data_bytes = open(os.path.join("examples", "student_scores.csv"), "rb").read()
    body = b""
    for field, fname, content, ctype in (
            ("paper", "paper.docx", paper_bytes,
             "application/vnd.openxmlformats-officedocument"
             ".wordprocessingml.document"),
            ("data", "data.csv", data_bytes, "text/csv")):
        body += (f"--{boundary}\r\n"
                 f'Content-Disposition: form-data; name="{field}"; '
                 f'filename="{fname}"\r\n'
                 f"Content-Type: {ctype}\r\n\r\n").encode() + content + b"\r\n"
    body += f"--{boundary}--\r\n".encode()
    req = urllib.request.Request(
        BASE + "/api/check_paper", data=body, method="POST",
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    with urllib.request.urlopen(req, timeout=120) as r:
        d = _json.loads(r.read())
    check("端到端 check_paper 成功", d.get("ok"))
    audit = d.get("audit", {})
    md_txt = audit.get("markdown", "") or ""
    check("端到端报告含表格核对小节", "表格交叉核对" in md_txt)
    check("端到端只报篡改那张（1 条 table_n）",
          sum(1 for c in audit.get("comparisons", [])
              if c.get("kind") == "table_n") == 1)
else:
    print("  [SKIP] 本地服务未启动（python app.py），端到端部分跳过")

print(f"\n{'=' * 50}")
print(f"表格交叉核查测试：{PASS} 通过 / {FAIL} 失败")
print(f"{'=' * 50}")
sys.exit(1 if FAIL else 0)
