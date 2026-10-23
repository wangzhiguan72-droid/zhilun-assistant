"""论文规范体检离线测试（v2.34）。全离线、零网络、零 LLM；不读仓库外文件，夹具全部现造。

覆盖：分节 / 文献条目解析 / GB/T 7714 单条规则 / 文献表汇总 / 正文引用 ↔ 文献表 /
图表编号 / 中文排版 / 结构 / 统计复算并入 / Markdown 与导出 / 红线措辞 / 边界与性能 /
docx 真文件链路 / /api/paper_lint / CLI lint。
每条误报护栏都来自真实论文调参时踩过的坑（页眉「参考文献」、全角空格、双语图题、续表……）。
"""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
os.environ.setdefault("RATE_LIMIT_DISABLE", "1")

from paper_lint import (_cn_int, _md_escape, check_entry, find_captions,  # noqa: E402
                        find_numeric_cites, lint_paper, parse_references, split_sections)
from stat_recheck import normalize  # noqa: E402

PASS = 0
FAIL = 0
Y = 2026   # 固定「今年」，测试不随系统日期漂移


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✅ {name}")
    else:
        FAIL += 1
        print(f"  ❌ {name} {detail}")


def lint(text, source="docx"):
    return lint_paper(text, source=source, current_year=Y)


def fs(res, cat=None, sub=None, pri=None):
    return [f for f in res["findings"] if (cat is None or f["category"] == cat)
            and (sub is None or sub in f["issue"]) and (pri is None or f["priority"] == pri)]


def refs_of(block, source="docx"):
    n = normalize(block)
    return parse_references(n, source, block)


print("[1] 分节：目录里的「参考文献」、页眉重复的「参考文献」、页码水印都不干扰")
DOC1 = ("目录\n第一章 绪论......1\n参考文献......20\n摘要\n本文研究数字金融[1]。\n"
        "第一章 绪论\n正文内容。\n参考文献\n"
        "[1] 张三. 数字金融与企业创新[J]. 经济研究, 2020, 55(3): 1-10.\n"
        "12\n参考文献\n万方数据\n"
        "[2] 李四. 产业经济学[M]. 北京: 高等教育出版社, 2019.\n"
        "致谢\n感谢导师。\n")
sec = split_sections(normalize(DOC1))
check("正文到真正的「参考文献」标题为止", DOC1[:sec["body"][1]].rstrip().endswith("正文内容。"), repr(DOC1[:sec["body"][1]][-20:]))
check("「致谢」之后归入文末附属部分", DOC1[sec["tail"][0]:].startswith("致谢"), repr(DOC1[sec["tail"][0]:][:10]))
p1 = parse_references(normalize(DOC1)[sec["refs"][0]:sec["refs"][1]], "docx")
check("页码 / 页眉 / 水印行不打断条目：解析出 [1][2]", [e["num"] for e in p1["entries"]] == [1, 2], str(p1))
check("条目正文里不混入页眉噪声", "万方" not in p1["entries"][0]["text"] and "参考文献" not in p1["entries"][0]["text"])
r1 = lint(DOC1)
check("整篇干净：0 条待核对", r1["counts"]["total"] == 0, str(r1["findings"]))
check("目录行不算章标题（章数 = 1）", r1["sections"]["chapters"] == 1, str(r1["sections"]))
check("无参考文献标题 → ref_heading False", split_sections("正文而已")["ref_heading"] is False)

print("\n[2] 文献条目解析：编号写法、全角空格、行内并条、PDF 断行、Word 自动编号")
p = refs_of("[1] 张三. 甲[J]. 乙, 2020, 1(1): 1-2.\n[2]　李四. 丙[J]. 丁, 2021, 2(1): 3-4.")
check("全角空格 / em-space 跟在 [n] 后照样切条", [e["num"] for e in p["entries"]] == [1, 2], str(p))
p = refs_of("[1] 张三. 甲[J]. 乙, 2020, 1(1): 1-2. [2] 李四. 丙[J]. 丁, 2021: 3-4.\n[3] 王五. 戊[M]. 北京: 某社, 2019.")
check("PDF 把 [2] 并到 [1] 行尾：按「上一条 +1」切开", [e["num"] for e in p["entries"]] == [1, 2, 3], str(p))
p = refs_of("[1] 张三. 数字金融. [9] 号文件的解读[J]. 乙, 2020, 1(1): 1-2.\n[2] 李四. 丙[J]. 丁, 2021, 2(1): 3-4.")
check("行内 [9] 不是「上一条 +1」，不切条", [e["num"] for e in p["entries"]] == [1, 2], str(p))
p = refs_of("1. 张三. 甲[J]. 乙, 2020, 1(1): 1-2.\n2. 李四. 丙[J]. 丁, 2021, 2(1): 3-4.")
check("「1.」式编号", [e["num"] for e in p["entries"]] == [1, 2] and p["mode"] == "explicit", str(p))
p = refs_of("[1] Smith J. Deep learn-\ning for finance[J]. Journal of\nFinance, 2020, 75(1): 1-30.\n[2] 张三. 数字金融与\n企业创新[J]. 经济研究, 2020, 55(3): 1-10.", "pdf")
check("PDF 断行：英文连字符拼回、英文补空格", "Deep learning for finance[J]. Journal of Finance" in p["entries"][0]["text"], p["entries"][0]["text"])
check("PDF 断行：中文直接相连", "数字金融与企业创新" in p["entries"][1]["text"], p["entries"][1]["text"])
IMPL = ("张三. 数字金融与企业创新[J]. 经济研究, 2020, 55(3): 1-10.\n"
        "李四. 产业经济学[M]. 北京: 高等教育出版社, 2019.\n"
        "王五. 供应链金融研究[D]. 上海: 复旦大学, 2021.")
p = refs_of(IMPL)
check("docx 丢了 Word 自动编号：按段落计为 1..3", p["mode"] == "implicit" and [e["num"] for e in p["entries"]] == [1, 2, 3], str(p))
check("同样文本来自 PDF：不猜编号（unparsed）", refs_of(IMPL, "pdf")["mode"] == "unparsed")
p = refs_of(IMPL + "\n变量 | 均值 | 标准差\n规模, 企业 | 1.2 | 0.3")
check("文末 docx 表格行不算文献条目", len(p["entries"]) == 3, str(len(p["entries"])))
check("空文献区 → none", refs_of("  \n ")["mode"] == "none")

print("\n[3] GB/T 7714 单条规则（每条都来自真实论文的写法）")


def ce(t):
    return check_entry(normalize(t), Y)


CASES = [
    ("规范期刊", "张三. 数字金融与企业创新[J]. 经济研究, 2020, 55(3): 1-10.", "J", []),
    ("只有期号「2020(04)」算有卷期", "张三. 标题[J]. 管理世界, 2020(04): 12-18.", "J", []),
    ("知网「2022,No.781(09)」", "李四. 标题[J]. 财会通讯, 2022,No.781(09): 1-3.", "J", []),
    ("网络首发「2025, 1-10.」算有页码", "王五. 标题[J]. 图书馆论坛, 2025, 1-10.", "J", []),
    ("外文文章号「130(3), 034504」", "Li X, Wang Y. Thin films[J]. Journal of Applied Physics, 2021, 130(3), 034504.", "J", []),
    ("期刊缺卷期和页码", "张三. 数字金融[J]. 经济研究, 2020.", "J", ["期刊缺卷(期)、页码"]),
    ("期刊缺页码（全角括号期号）", "李红平. 供应链管理模式下的采购管理[J]. 中国科技投资, 2021（5）.", "J", ["期刊缺页码"]),
    ("PDF 把 2016 拆成「20 16」", "赵六. 标题[J]. 会计研究, 20 16, 5(2): 1-9.", "J", []),
    ("出版年晚于今年", "赵六. 标题[J]. 会计研究, 2031, 5(2): 1-9.", "J", ["出版年晚于今年"]),
    ("[J/OL] 免卷期页码，但缺网址", "王秀红, 吴丹, 刘浩东, 等. 知识图谱构建[J/OL]. 图书馆杂志, 1-20[2025-09-28].",
     "J/OL", ["网络首发文献缺网址 / DOI"]),
    ("[J/OL] 有 DOI 即可", "王秀红. 知识图谱构建[J/OL]. 图书馆杂志, 1-20[2025-09-28]. DOI: 10.13663/j.cnki.lj.2025.01.001.",
     "J/OL", []),
    ("[EB/OL] 缺网址、引用日期（圆括号是发布日期）", "国家统计局. 2023年统计公报[EB/OL]. (2024-02-28).",
     "EB/OL", ["网络资源缺网址 / DOI", "电子文献缺引用日期 [YYYY-MM-DD]"]),
    ("[EB/OL] 齐全", "国家统计局. 2023年统计公报[EB/OL]. (2024-02-28)[2026-09-01]. https://www.stats.gov.cn/.", "EB/OL", []),
    ("专著缺出版项", "张三. 产业经济学[M]. 2019.", "M", ["专著缺「出版地: 出版者, 年」"]),
    ("专著齐全", "张三. 产业经济学[M]. 北京: 高等教育出版社, 2019.", "M", []),
    ("专著缺年：只报缺年，不再叠报出版项", "马克思, 恩格斯. 马克思恩格斯选集[M]. 北京: 人民出版社.", "M", ["缺出版年"]),
    ("学位论文缺保存地", "陈成天. 支持中小企业融资的财政政策研究[D]. 财政部财政科学研究所, 2015: 35-43.", "D",
     ["学位论文缺「保存地: 保存单位, 年」"]),
    ("学位论文齐全", "陈成天. 财政政策研究[D]. 北京: 财政部财政科学研究所, 2015.", "D", []),
    ("标准 [S] 可不写年", "全国信息与文献标准化技术委员会. 信息与文献 参考文献著录规则: GB/T 7714[S].", "S", []),
    ("无类型码：只报给汇总层，单条不叠加期刊规则", "Abualigah L. 2025. Particle swarm optimization. Archives, 32(1): 1-20.", None, []),
    ("作者 4 人未加「等」", "张三, 李四, 王五, 赵六. 标题[J]. 经济研究, 2020, 55(3): 1-10.", "J",
     ["作者超过 3 人未写「等 / et al.」"]),
    ("作者 3 人 + 等", "张三, 李四, 王五, 等. 标题[J]. 经济研究, 2020, 55(3): 1-10.", "J", []),
    ("外文作者 4 人未加 et al.", "Smith J, Brown K, Lee M, Wang X. Title[J]. Nature, 2020, 580(7801): 1-5.", "J",
     ["作者超过 3 人未写「等 / et al.」"]),
    ("机构作者不数人头", "中国人民银行营业管理部课题组. 标题[R]. 北京: 中国人民银行, 2020.", "R", []),
]
for name, text, code, probs in CASES:
    got_code, got = ce(text)
    check(name, got_code == code and got == probs, f"→ {got_code} {got}")

print("\n[4] 文献表汇总：缺类型标识、疑似重复、编号跳号 / 重号、同类问题聚合")


def with_refs(refs, body="摘要\n本文研究数字金融。\n"):
    return body + "参考文献\n" + refs


r = lint(with_refs("[1] 张三. 甲[J]. 乙, 2020, 1(1): 1-2.\n[2] 李四. 丙. 丁, 2021, 2(1): 3-4.\n"
                   "[3] 王五. 戊[M]. 北京: 某出版社, 2019."))
check("部分条目缺类型标识 → P1 点名 [2]", any("[2]" in f["issue"] for f in fs(r, "参考文献", "缺文献类型标识", "P1")), str(r["findings"]))
r = lint(with_refs("[1] Smith J. 2020. Title one. Journal A, 1(1): 1-2.\n[2] Brown K. 2021. Title two. Journal B, 2(1): 3-4."))
f = fs(r, "参考文献", "全部 2 条文献都没有文献类型标识")
check("全部无类型标识（APA 式）→ 汇总成一条，提示可按期刊要求忽略", len(f) == 1 and "APA" in f[0]["action"], str(r["findings"]))
r = lint(with_refs("[1] 张三. 数字金融与企业创新[J]. 经济研究, 2020, 55(3): 1-10.\n"
                   "[2] 李四. 产业经济学[M]. 北京: 高等教育出版社, 2019.\n"
                   "[3] 张三.数字金融与企业创新 [J].经济研究,2020,55(3):1-10."))
check("同一文献列两次（空格标点不同）→ 疑似重复 [1] 与 [3]", bool(fs(r, "参考文献", "[1] 与 [3]", "P1")), str(r["findings"]))
r = lint(with_refs("[1] 张三. 甲[J]. 乙, 2020, 1(1): 1-2.\n[2] 李四. 丙[J]. 丁, 2021, 2(1): 3-4.\n"
                   "[4] 王五. 戊[J]. 己, 2022, 3(1): 5-6.\n[4] 赵六. 庚[J]. 辛, 2023, 4(1): 7-8."))
check("编号跳号 + 重号 → 一条 P1 同时点名", bool(fs(r, "参考文献", "缺 [3]；重复 [4]", "P1")), str(r["findings"]))
many = "\n".join(f"[{i}] 作者{chr(0x4e00 + i * 7)}. 研究题目第{i}篇[J]. 期刊名, 2020." for i in range(1, 26))
r = lint(with_refs(many))
f = fs(r, "参考文献", "期刊缺卷(期)、页码")
check("25 条同类问题聚合成 1 条，列前 20 条 + 「等 25 条」", len(f) == 1 and "等 25 条" in f[0]["issue"], str([x["issue"] for x in f]))
check("汇总信息：25 条、按编号解析", r["sections"]["references"] == 25 and r["sections"]["ref_mode"] == "explicit")
check("有文献表但正文无引用 → P2 提示（上标引用抽取会丢）", bool(fs(r, "引用核对", "没识别到", "P2")), str(r["findings"]))

print("\n[5] 正文引用 ↔ 文献表")
cites = find_numeric_cites(normalize("研究表明[1]，另见[2,5]与[3-4]；[6]-[8] 讨论了机制，Smith[9] 提出了模型。"
                                     "取值区间[0,1]内，数组 a[1]、x[2] 与区间[2,5]不是引用。"))
flat = [n for _, _, ns in cites for n in ns]
check("[1] [2,5] [3-4] [6]-[8] Smith[9] 都认；[0,1] a[1] x[2] 区间[2,5] 不认", flat == [1, 2, 5, 3, 4, 6, 7, 8, 9], str(flat))
NUM_REFS = ("参考文献\n"
            "[1] 张三, 李四, 王五, 赵六. 数字金融与企业创新[J]. 经济研究, 2020, 55(3): 1-10.\n"
            "[2] 钱七. 融资约束研究[J]. 管理世界, 2019, 35(2).\n"
            "[3] 孙八. 产业经济学[M]. 北京: 高等教育出版社, 2018.\n"
            "[4] 周九. 供应链金融研究[D]. 上海: 复旦大学, 2021.\n"
            "[5] 吴十. 中小企业融资报告[R]. 北京: 某研究院, 2022.\n"
            "[6] 郑十一. 未被引用的文献[J]. 会计研究, 2017, 12(1): 5-9.\n")
NUM_BODY = ("摘要\n本文研究数字金融对企业创新的影响。\n关键词：数字金融；企业创新；融资约束\n"
            "第一章 绪论\n数字金融近年来发展迅速[1]。\n"
            "第二章 文献综述\n已有文献讨论了融资约束[2,3]，也有学者关注中小企业[5]与供应链[4]。\n"
            "第三章 实证分析\n回归结果显示政策效果显著，t(58) = 2.35, p = .30，说明政策有效。\n")
r = lint(NUM_BODY + NUM_REFS)
check("引用体例 = 顺序编码制，4 处引用", r["sections"]["cite_style"] == "numeric" and r["sections"]["citations"] == 4, str(r["sections"]))
check("未被引用 [6] → P1", bool(fs(r, "引用核对", "[6] 在正文里没有被引用", "P1")), str(r["findings"]))
check("[4] 晚于 [5] 首次出现 → 顺序 P2", bool(fs(r, "引用核对", "[4] 在 [5] 之后", "P2")), str(r["findings"]))
check("无越界引用", not fs(r, "引用核对", "文献表里没有对应编号"))
check("[2] 期刊缺页码 → P1", bool(fs(r, "参考文献", "[2] 期刊缺页码", "P1")), str(r["findings"]))
check("[1] 4 位作者未加「等」→ P2", bool(fs(r, "参考文献", "[1] 作者超过 3 人", "P2")))
check("统计复算并入：t(58)=2.35 却写 p=.30 → P0 决策错误", bool(fs(r, "统计复算", pri="P0")), str(r["findings"]))
check("结构完整：关键词 3 个、章 3 个", r["sections"]["keywords"] == 3 and r["sections"]["chapters"] == 3, str(r["sections"]))
check("排序：P0 在最前", r["findings"][0]["priority"] == "P0")
r = lint(NUM_BODY + "此外有研究指出[8]。\n" + NUM_REFS)
f = fs(r, "引用核对", "[8]，文献表里没有对应编号", "P0")
check("越界引用 [8] → P0，位置摘录可 Ctrl+F", len(f) == 1 and "[8]" in f[0]["context"], str(r["findings"]))
impl_refs = "参考文献\n" + "\n".join(ln.split("] ", 1)[1] for ln in NUM_REFS.splitlines()[1:])
r = lint(NUM_BODY + "此外有研究指出[8]。\n" + impl_refs)
check("Word 自动编号（按段落数）", r["sections"]["ref_mode"] == "implicit" and r["sections"]["references"] == 6, str(r["sections"]))
check("按段落计数时越界降为 P1 并注明口径", bool(fs(r, "引用核对", "按文献表段落顺序计 6 条", "P1")), str(r["findings"]))
check("按段落计数时不报首次引用顺序", not fs(r, "引用核对", "首次引用顺序"))
r = lint(NUM_BODY + impl_refs, source="pdf")
check("PDF + 无编号文献表 → 不猜，P2 说明跳过", bool(fs(r, "引用核对", "无法解析", "P2")), str(r["findings"]))
r = lint(NUM_BODY + "此外有研究指出[8]。\n" + NUM_REFS, source="pdf")
check("PDF：越界 P0 → P1、未引用 P1 → P2", bool(fs(r, "引用核对", "[8]", "P1")) and bool(fs(r, "引用核对", "[6]", "P2")),
      str(r["findings"]))
check("PDF：统计复算不降级（数字抽取可靠）", bool(fs(r, "统计复算", pri="P0")))
check("PDF：附说明", any("PDF" in n for n in r["notes"]))
LEAK = ("摘要\n本文研究数字金融。\n已有研究[1]、[2]、[3]讨论了融资约束，也有研究[12]关注创新。\n"
        "[12] 钱七. 被双栏排版挤进正文的文献[J]. 管理世界, 2019, 35(2): 1-9.\n"
        "参考文献\n" + "\n".join(f"[{i}] 作者{chr(0x4e00 + i * 11)}. 题目第{i}篇[J]. 期刊, 2020, 1(1): 1-2." for i in range(1, 12)))
r = lint(LEAK, source="pdf")
check("双栏 PDF 把 [12] 条目排进正文：不算越界引用", not fs(r, "引用核对", "文献表里没有对应编号"), str(r["findings"]))
AY = ("摘要\n本文研究数字金融。\n"
      "已有研究表明数字金融显著促进了创新（张三，2020；李四等，2019），"
      "但也有研究得出相反结论（王五，2021）。国外学者（Smith et al., 2018）与（Brown & Lee, 2017）也有讨论。\n"
      "参考文献\n"
      "张三. 数字金融与企业创新[J]. 经济研究, 2020, 55(3): 1-10.\n"
      "李四, 钱七, 孙八. 融资约束与创新[J]. 管理世界, 2019, 35(2): 1-9.\n"
      "Smith J, Brown K, Lee M, et al. Fintech and innovation[J]. Journal of Finance, 2018, 73(1): 1-30.\n")
r = lint(AY)
check("著者-出版年制识别", r["sections"]["cite_style"] == "author_year", str(r["sections"]))
f = fs(r, "引用核对", "有 2 处著者-出版年引用在文献表里找不到", "P1")
check("（王五，2021）（Brown & Lee, 2017）找不到 → P1；其余三处匹配", len(f) == 1 and "王五" in f[0]["issue"]
      and "Brown" in f[0]["issue"], str(r["findings"]))
r = lint("摘要\n本文研究。\n已有研究[1]、[2]、[3]讨论了该问题。\n文献\n张三. 甲[J]. 乙, 2020.\n")
check("有 3 处引用却没有「参考文献」标题 → P1", bool(fs(r, "引用核对", "没找到「参考文献」一节", "P1")), str(r["findings"]))

print("\n[6] 图表编号")
FIG = ("如图1所示，由图2可知，如图5所示，另见图6。\n"
       "图1 研究框架\n图2 技术路线\n图4 样本分布\n图6 结果对比\n"
       "表1中可以看出，收入逐年增加，见表1与表2。\n"
       "表1 描述性统计\n表1 回归结果\n表2 稳健性检验\n")
r = lint(FIG)
check("识别 4 个图题、3 个表题（「表1中可以看出」是正文不是表题）",
      r["sections"]["figures"] == 4 and r["sections"]["tables"] == 3, str(r["sections"]))
check("图号跳号 → P2 缺图3、图5", bool(fs(r, "图表编号", "缺 图 3、图 5", "P2")), str(r["findings"]))
check("正文提到图5但没有图题 → P1", bool(fs(r, "图表编号", "提到了 图 5", "P1")), str(r["findings"]))
check("图4 未被正文提及 → P2", bool(fs(r, "图表编号", "图 4 在正文里没有被提及", "P2")), str(r["findings"]))
check("表1 对应两个标题 → 重号 P1", bool(fs(r, "图表编号", "表 1 对应了不同标题", "P1")), str(r["findings"]))
FIG2 = ("结构如图1-1、图1-2和图2-1所示，结果见表4-1、表4-2和表4-3。\n"
        "图1-1 桥式起重机      图1-2 门式起重机\n"
        "Fig. 1-1 Bridge crane      Fig. 1-2 Gantry crane\n"
        "图2-1 控制框图\nFig. 2-1 Control diagram\n"
        "表4-1 变量定义\n表4-2-1 2008-2014 年描述性统计\n表4-2-2 2015-2021 年描述性统计\n"
        "表4-3 回归结果\n表4-3(续) 回归结果\n")
r = lint(FIG2)
check("中英双语图题不算重号；同一行并排两图都识别", r["sections"]["figures"] == 3, str(r["sections"]))
check("拆分子表 4-2-1/4-2-2 算「表 4-2」在；续表不算重号", r["sections"]["tables"] == 4 and not fs(r, "图表编号"),
      str(r["findings"]))
r = lint("三张图的关系如图1-3所示。\n图1 甲\n图2 乙\n图3 丙\n")
check("平铺编号下「图1-3」是范围：图1..3 都算被提及", not fs(r, "图表编号"), str(r["findings"]))
r = lint("如图1所示，另见图2。\n正文没有任何图题。\n")
check("提到图号却一个图题都没识别到 → P2（图题可能在文本框里）", bool(fs(r, "图表编号", "没识别到图题", "P2")), str(r["findings"]))
r = lint("本文分析表1的结果，并与已有研究比较。\n表1 描述性统计\n")
check("「分析表1」宽松提及：不报未提及", not fs(r, "图表编号", "没有被提及"), str(r["findings"]))
r = lint("量表共分三个维度，每个维度代表3类能力，作者发表2篇论文。\n表1 描述性统计\n见表1。\n")
check("「代表3」「发表2」不是表号：不报「提到了表 3」", not fs(r, "图表编号", "提到了"), str(r["findings"]))
r = lint(FIG, source="pdf")
check("PDF：图表 P1 降为 P2", not fs(r, "图表编号", pri="P1") and bool(fs(r, "图表编号", "提到了 图 5", "P2")),
      str(r["findings"]))

print("\n[7] 中文排版（docx）")


def typo(line):
    return fs(lint("摘要\n" + line + "\n"), "中文排版")


def typo_hit(line, name):
    return any(f["issue"].startswith(name) for f in typo(line))


check("中文之间半角逗号 / 分号 → 2 处", any(f["issue"] == "中文之间用了半角标点：2 处"
                                   for f in typo("学生普遍认为该方法有效,但也有人提出不同意见;需要进一步讨论研究。")))
check("半角括号括中文", typo_hit("问卷发放采用两种途径(包括线上与线下两种方式)，共回收问卷三百份。", "半角括号括中文"))
check("半角括号括英文缩写不报", not typo_hit("本研究的样本数据来源于国泰安数据库(CSMAR)与万得数据库。", "半角括号括中文"))
check("重复标点 → 2 处", any(f["issue"] == "重复标点：2 处" for f in typo("研究发现，，数字金融显著促进了企业创新。。这一结论较为稳健。")))
check("「......」省略号不规范", typo_hit("他说到：“大嫂，上头下来人......说是又要征兵嘞，大家都得去。”", "省略号写法不规范"))
check("「……」规范写法不报", not typo_hit("他说到：“大嫂，上头下来人……说是又要征兵嘞，大家都得去。”", "省略号写法不规范"))
check("全角数字 → 2 处", any(f["issue"] == "全角数字：2 处" for f in typo("本次问卷共发放３００份，回收有效问卷２８０份，有效率较高。")))
check("中文间多余空格 ≥3 处才报", typo_hit("数字 金融 对 企业 创新 的 影响 研究 综述 与 展望 分析", "中文之间有多余空格"))
check("零星 1 处空格不打扰", not typo_hit("本文研究数字金融对企业 创新的影响机制与作用路径分析。", "中文之间有多余空格"))
check("书名号不成对", typo_hit("根据《中华人民共和国统计法规定，本研究的数据均经过脱敏处理。", "括号 / 引号不成对"))
check("跨段引文（段首有前引号、段末不收）不报", not typo_hit("“这是一段很长的跨段引文，引文的第一段到这里结束，下段继续。", "括号 / 引号不成对"))
check("docx 表格行不查排版", not typo("变量, 名称 | 均值, 标准差 | 最小值, 最大值"))
check("20 字以下短行不查", not typo("研究,方法"))
r = lint("摘要\n本文研究数字金融。\n参考文献\n[1] 张三,李四.数字金融与企业创新[J].经济研究,2020,55(3):1-10.\n")
check("参考文献区不查中文排版（知网导出本就是半角）", not fs(r, "中文排版"), str(r["findings"]))
r = lint("摘要\n学生普遍认为该方法有效,但也有人提出不同意见;需要进一步讨论研究。\n", source="pdf")
check("PDF 跳过中文排版并说明", not fs(r, "中文排版") and any("中文排版项已跳过" in n for n in r["notes"]))

print("\n[8] 结构：关键词、中英文摘要、章编号")
check("关键词 2 个 → P2", bool(fs(lint("摘要\n正文。\n关键词：数字金融；企业创新\n"), "论文结构", "关键词识别到 2 个", "P2")))
kw9 = "；".join(["甲乙", "丙丁", "戊己", "庚辛", "壬癸", "子丑", "寅卯", "辰巳", "午未"])
check("关键词 9 个 → P2", bool(fs(lint("摘要\n正文。\n关键词：" + kw9 + "\n"), "论文结构", "关键词识别到 9 个")))
r = lint("摘要\n正文。\n关键词：存货管理 供应链管理 光明乳业\n")
check("单空格分隔的 3 个关键词不误报", r["sections"]["keywords"] == 3 and not fs(r, "论文结构"), str(r["sections"]))
check("PDF 断行只剩 2 个关键词：不报", not fs(lint("摘要\n正文。\n关键词：数字金融；企业创新\n", "pdf"), "论文结构"))
FILL = "这是用于测试篇幅的正文内容。" * 1200 + "\n"
r = lint("摘要\n本文研究数字金融。\n关键词：甲乙；丙丁；戊己\n" + FILL)
check("长文有中文摘要无英文摘要 → P2", bool(fs(r, "论文结构", "英文摘要", "P2")), str(r["findings"]))
r = lint("摘要\n本文研究数字金融。\n关键词：甲乙；丙丁；戊己\nAbstract\nThis paper studies fintech.\n" + FILL)
check("有 Abstract 不报", not fs(r, "论文结构", "英文摘要"))
r = lint(FILL + "参考文献\n[1] 张三. 甲[J]. 乙, 2020, 1(1): 1-2.\n")
check("长文 + 有文献表却没有摘要 → P1", bool(fs(r, "论文结构", "没识别到「摘要」", "P1")), str(r["findings"]))
CH = ("目录\n第三章 研究设计......12\n摘要\n正文。\n第一章 绪论\n正文。\n第二章 文献综述\n正文。\n"
      "第四章 结论\n正文。\n")
r = lint(CH)
check("章编号 1、2、4 → 缺第3章 P1（目录里的第三章不算）", bool(fs(r, "论文结构", "缺 第3章", "P1")), str(r["findings"]))
check("中文数字章号换算", [_cn_int(s) for s in ("三", "十", "十二", "二十", "二十一", "7")] == [3, 10, 12, 20, 21, 7])

print("\n[9] 统计复算并入：只复算正文与附录，不复算文献标题")
r = lint("摘要\n本文研究。\n参考文献\n[1] 张三. 效应检验 t(58) = 2.35, p = .30 的再讨论[J]. 乙, 2020, 1(1): 1-2.\n")
check("文献标题里的检验式不复算", r["stats"]["summary"]["checked"] == 0, str(r["stats"]["summary"]))
r = lint("摘要\n本文研究。\n参考文献\n[1] 张三. 甲[J]. 乙, 2020, 1(1): 1-2.\n附录\n补充检验：t(58) = 2.35, p = .022。\n")
check("附录里的检验式照常复算", r["stats"]["summary"]["checked"] == 1 and not fs(r, "统计复算"), str(r["stats"]["summary"]))
check("stats 字段齐全", all(k in r["stats"] for k in ("alpha", "summary", "summary_text", "items", "method_note", "disclaimer")))

print("\n[10] Markdown 报告、Word 导出、红线措辞")
r = lint(NUM_BODY + "此外有研究指出[8]。\n" + NUM_REFS)
md = r["markdown"]
check("报告标题与分级小节", md.startswith("# 论文提交前规范体检") and "## 必改" in md and "## 建议改" in md and "## 可选" in md)
check("每条问题都有「建议」行", md.count("  - 建议：") == r["counts"]["total"], f"{md.count('  - 建议：')} vs {r['counts']['total']}")
check("计数自洽", r["counts"]["total"] == len(r["findings"]) == r["counts"]["P0"] + r["counts"]["P1"] + r["counts"]["P2"]
      == sum(r["by_category"].values()))
check("按优先级排序", [f["priority"] for f in r["findings"]] == sorted(f["priority"] for f in r["findings"]))
check("字段齐全", all(set(f) >= {"priority", "category", "issue", "action", "context"} for f in r["findings"]))
check("Markdown 转义：拆 ** 并压成单行", _md_escape("a**b**\n  c") == "a*b* c")
from export_docx import markdown_to_docx  # noqa: E402
blob = markdown_to_docx(md, title="论文规范体检")
check("走 /api/export 同一函数能导出 Word", blob[:2] == b"PK" and len(blob) > 5000, str(len(blob)))
from docx import Document  # noqa: E402
doc_text = "\n".join(p.text for p in Document(io.BytesIO(blob)).paragraphs)
check("导出的 Word 里有结论和必改项", "论文提交前规范体检" in doc_text and "文献表里没有对应编号" in doc_text)
ALL = [lint(t) for t in (DOC1, NUM_BODY + NUM_REFS, AY, FIG, FIG2, CH, LEAK)]
blob_txt = json.dumps(ALL, ensure_ascii=False)
check("任何输出都不出现「造假 / 作弊 / 伪造」", not any(w in blob_txt for w in ("造假", "作弊", "伪造")))
check("统计复算声明「不构成任何学术不端判定」随报告输出", "不构成任何学术不端判定" in md)
check("规范体检声明「不改写正文」", "不改写正文" in r["disclaimer"])

print("\n[11] 边界与性能")
for empty in ("", "   \n\n", None):
    r = lint(empty)
    check(f"空输入 {empty!r} 不崩、0 条", r["ok"] and r["counts"]["total"] == 0 and "未发现" in r["headline"])
refs = "\n".join(f"[{i}] 作者{chr(0x4e00 + i)}{chr(0x4e00 + i + 400)}. 研究题目第{i}篇[J]. 期刊名, 2020, {i}(1): {i}-{i + 5}."
                 for i in range(1, 301))
body = "摘要\n本文研究。\n关键词：甲乙；丙丁；戊己\n" + "".join(
    f"第{i}段研究发现[{i}]，如图{(i - 1) % 80 + 1}所示，t(58) = 2.35, p = .022。\n" for i in range(1, 301))
body += "".join(f"图{k} 第{k}幅图\n" for k in range(1, 81))
t0 = time.time()
r = lint(body + "参考文献\n" + refs)
dt = time.time() - t0
check(f"300 条文献 + 300 处引用 + 80 个图 + 300 处检验 < 8s（{dt:.2f}s）", dt < 8)
check("大文档无误报", r["counts"]["total"] == 0, str(r["findings"][:3]))
check("大文档统计复算 300 处全部一致", r["stats"]["summary"]["checked"] == 300 and r["stats"]["summary"]["consistent"] == 300,
      str(r["stats"]["summary"]))
EVIL = ["." * 200000, "中" + "…" * 100000, "如图1" * 60000, "[1]" * 80000, "(" + "中" * 100000,
        "参考文献\n[1] 张三. 甲[J]. 乙, 2020(" + "1 " * 50000, ("目录" + "." * 5000 + "x\n") * 200]
t0 = time.time()
for t in EVIL:
    lint("摘要\n" + t + "\n")
dt = time.time() - t0
check(f"对抗输入（超长点线 / 省略号 / 括号 / 卷期）无回溯爆炸 < 10s（{dt:.2f}s）", dt < 10)

print("\n[12] docx 真文件链路（read_paper_text 丢空段、表格行附在文末）")


def make_docx(paras, table=None):
    d = Document()
    for ptxt in paras:
        d.add_paragraph(ptxt)
        d.add_paragraph("")                         # 空段会被 read_paper_text 丢掉
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


DOCX = make_docx(NUM_BODY.strip().split("\n") + ["此外有研究指出[8]。"] + NUM_REFS.strip().split("\n"),
                 table=[["变量", "均值", "标准差"], ["规模,企业", "1.20", "0.30"]])


class _FS:
    def __init__(self, name, data):
        self.filename, self._d = name, data

    def read(self):
        return self._d


from extract_paper import read_paper_text  # noqa: E402
dtext = read_paper_text(_FS("论文.docx", DOCX))
r = lint(dtext, "docx")
check("docx：6 条文献按编号解析", r["sections"]["references"] == 6 and r["sections"]["ref_mode"] == "explicit", str(r["sections"]))
check("docx：越界 [8] 仍是 P0", bool(fs(r, "引用核对", "[8]", "P0")))
check("docx：表格行不进文献表、不查排版", not fs(r, "中文排版") and not fs(r, "参考文献", "变量"), str(r["findings"]))

print("\n[13] /api/paper_lint（Flask test_client）")
import app as flask_app  # noqa: E402
client = flask_app.app.test_client()
resp = client.post("/api/paper_lint", json={"text": NUM_BODY + "此外有研究指出[8]。\n" + NUM_REFS})
j = resp.get_json()
check("JSON 文本 → 200 + 报告", resp.status_code == 200 and j["ok"] and j["counts"]["P0"] >= 2, str(j)[:300])
check("返回 markdown（前端直接走 /api/export）与 text_length", j["markdown"].startswith("# 论文提交前规范体检") and j["text_length"] > 0)
resp = client.post("/api/paper_lint", json={"text": "  "})
check("空文本 → 400", resp.status_code == 400 and not resp.get_json()["ok"])
resp = client.post("/api/paper_lint", data="[1,2]", content_type="application/json")
check("非对象 JSON → 400 而不是 500", resp.status_code == 400, str(resp.status_code))
resp = client.post("/api/paper_lint", json={"text": "摘要\n学生普遍认为该方法有效,但也有人提出不同意见;需要进一步讨论研究。\n",
                                            "source": "pdf"})
check("source=pdf 生效（跳过排版并说明）", any("PDF" in n for n in resp.get_json()["notes"]))
resp = client.post("/api/paper_lint", json={"text": "摘要\n正文内容。\n", "source": "../../etc"})
check("非法 source 回落为 docx", resp.status_code == 200 and resp.get_json()["source"] == "docx")
resp = client.post("/api/paper_lint", data={"paper": (io.BytesIO(DOCX), "论文.docx")}, content_type="multipart/form-data")
j = resp.get_json()
check("multipart 上传 docx → 200，6 条文献", resp.status_code == 200 and j["sections"]["references"] == 6, str(j)[:300])
resp = client.post("/api/paper_lint", data={"paper": (io.BytesIO(b"x"), "a.exe")}, content_type="multipart/form-data")
check("不支持的扩展名 → 400", resp.status_code == 400)
resp = client.post("/api/paper_lint", json={"text": "中" * (flask_app.PAPER_LINT_MAX_CHARS + 1)})
check("超长文本 → 400（不进规则引擎）", resp.status_code in (400, 413), str(resp.status_code))

print("\n[14] CLI：python cli.py lint")
env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
with tempfile.TemporaryDirectory() as td:
    bad = Path(td) / "论文.txt"
    bad.write_text(NUM_BODY + "此外有研究指出[8]。\n" + NUM_REFS, encoding="utf-8")
    good = Path(td) / "干净.txt"
    good.write_text(DOC1, encoding="utf-8")


    def run_cli(*argv):
        pr = subprocess.run([sys.executable, str(ROOT / "cli.py"), "lint", *argv], cwd=str(ROOT), env=env,
                            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)
        return pr.returncode, pr.stdout, pr.stderr


    rc, out, err = run_cli(str(bad))
    check("有必改项 → 退出码 1，stdout 是 Markdown", rc == 1 and out.startswith("# 论文提交前规范体检"), f"rc={rc} {err[:200]}")
    rc, out, err = run_cli(str(bad), "-f", "json", "--exit-zero")
    try:
        jj = json.loads(out)
    except ValueError:
        jj = {}
    check("--format json --exit-zero → 0，JSON 不含 markdown", rc == 0 and jj.get("counts", {}).get("P0", 0) >= 2
          and "markdown" not in jj, f"rc={rc} {err[:200]}")
    rc, out, err = run_cli(str(good))
    check("干净论文 → 退出码 0", rc == 0, f"rc={rc} {err[:200]}")
    outp = Path(td) / "体检.md"
    rc, out, err = run_cli(str(bad), "-o", str(outp), "--exit-zero")
    check("-o 写文件", rc == 0 and outp.read_text(encoding="utf-8").startswith("# 论文提交前规范体检"), f"rc={rc}")
    rc, out, err = run_cli(str(Path(td) / "不存在.docx"))
    check("文件不存在 → 退出码 2 + 可读提示", rc == 2 and "不存在" in err, f"rc={rc} {err[:200]}")

print(f"\n结果:{PASS} 通过 / {FAIL} 失败")
sys.exit(1 if FAIL else 0)
