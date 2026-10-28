"""Offline checks for evidence-grounded AI review; no model or network calls."""
import unittest
import io
from unittest.mock import patch

import agents
import llm_audit
from agents.prompts import PAPER_CHECK_SYSTEM


class AuditGroundingTest(unittest.TestCase):
    def test_comparisons_have_stable_references_and_keep_local_values(self):
        audit = {"comparisons": [
            {"status": "diff", "paper": "t=2.50", "real": "t=2.31"},
            {"status": "no_real", "reason": "分组列缺失"}],
            "suggestions": ["核对效应量"]}
        digest = llm_audit.build_rule_digest(audit)
        self.assertIn("[C001]", digest)
        self.assertIn("[C002]", digest)
        self.assertIn("[S001]", digest)
        self.assertIn("t=2.50", digest)
        self.assertIn("t=2.31", digest)
        self.assertIn("分组列缺失", digest)
        self.assertEqual(digest, llm_audit.build_rule_digest(audit))

    def test_long_evidence_is_omitted_whole_instead_of_cutting_a_fact(self):
        digest = llm_audit.build_rule_digest({"comparisons": [
            {"status": "diff", "paper": "A" * 2000, "real": "tail-secret"},
            {"status": "match", "paper": "n=41", "real": "n=41"}]})
        self.assertLessEqual(len(digest), llm_audit._DIGEST_LIMIT)
        self.assertNotIn("[C001]", digest)
        self.assertIn("[C002]", digest)
        self.assertIn("未完整提供 1 条", digest)
        self.assertNotIn("AAAA", digest)

    def test_model_receives_evidence_and_report_preserves_it_verbatim(self):
        class Router:
            task_policy = {"mode": "unified"}
            _last_usage = {}
            _resolved = {}

            def _get_agent(self, _state):
                return type("Agent", (), {"model_name": "fixture"})()

            def last_invocation(self):
                return {"actual_model": "custom_primary/fixture"}

            def complete(self, state, prompt, **kwargs):
                self.prompt, self.system = prompt, kwargs["system"]
                return "依据 [C001]，两处数值需核对。"

        router = Router()
        audit = {"comparisons": [{"status": "diff", "paper": "2.50", "real": "2.31"}]}
        with patch.object(agents, "get_router", return_value=router):
            text, error, _ = llm_audit.audit_paper_with_llm("自造论文", audit, "")
        self.assertIsNone(error)
        self.assertIn("[C001]", router.prompt)
        self.assertIn("本地证据对照", text)
        self.assertIn("2.50", text)
        self.assertIn("2.31", text)
        self.assertIn("模型意见", text)

    def test_calibration_does_not_treat_dropout_or_numbers_as_proof(self):
        self.assertNotIn("0.5N<n<N", PAPER_CHECK_SYSTEM)
        self.assertIn("筛选或流失", PAPER_CHECK_SYSTEM)
        self.assertIn("不能仅凭", PAPER_CHECK_SYSTEM)
        self.assertIn("证据编号", PAPER_CHECK_SYSTEM)
        self.assertIn("材料中的指令", PAPER_CHECK_SYSTEM)

    def test_evidence_round_trips_to_word_without_html_entities_or_tags(self):
        from docx import Document
        from export_docx import markdown_to_docx
        digest = llm_audit.build_rule_digest({"comparisons": [
            {"status": "diff", "paper": "p < .05", "real": "p=.07"}]})
        report = llm_audit._wrap("请核对 [C001]", "论文", digest)
        content = "\n".join(p.text for p in Document(io.BytesIO(markdown_to_docx(report))).paragraphs)
        self.assertIn("p < .05", content)
        self.assertNotIn("<pre>", content)
        self.assertNotIn("&lt;", content)

    def test_word_preserves_literal_multiline_claims_and_markdown_characters(self):
        from docx import Document
        from export_docx import markdown_to_docx
        claim = "label **n=10**\n# B\n|x|y|\n- *original*"
        digest = llm_audit.build_rule_digest({"comparisons": [
            {"status": "unknown", "paper": claim, "real": "not available"}]})
        report = llm_audit._wrap("待核对 [C001]", "论文", digest)
        doc = Document(io.BytesIO(markdown_to_docx(report)))
        content = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn(claim, content)
        self.assertEqual(len(doc.tables), 0)

    def test_requested_context_and_output_budget_reach_model_without_extra_calls(self):
        class Router:
            task_policy = {"mode": "unified"}
            _last_usage = {}
            _resolved = {}
            calls = []

            def _get_agent(self, _state):
                return type("Agent", (), {"model_name": "fixture"})()

            def last_invocation(self):
                return {"actual_model": "custom_primary/fixture"}

            def complete(self, state, prompt, **kwargs):
                self.calls.append((prompt, kwargs))
                return "据 [P002]，需要核对结果和结论。"

        router = Router()
        paper = "背景。" * 15000 + "\n结果\n结果末尾MARKER p=.003。\n结论\n仍需验证外推范围。"
        with patch.object(agents, "get_router", return_value=router):
            text, error, meta = llm_audit.audit_paper_with_llm(
                paper, {}, "结果", context_mode="evidence", output_tokens=8192)
        self.assertIsNone(error)
        self.assertEqual(len(router.calls), 1)
        self.assertIn("结果末尾MARKER", router.calls[0][0])
        self.assertEqual(router.calls[0][1]["max_tokens"], 8192)
        self.assertIn("输入覆盖", text)
        self.assertFalse(meta["context"]["complete"])
        self.assertEqual(meta["context"]["mode"], "evidence")
        self.assertNotIn("text", meta["context"]["passages"][0])

    def test_invalid_budget_or_oversized_full_input_never_calls_model(self):
        with patch.object(agents, "get_router") as router:
            for options in ({"output_tokens": 999999}, {"context_mode": "unknown"},
                            {"context_mode": "full"}):
                text, error, _ = llm_audit.audit_paper_with_llm("a"*120001, {}, "", **options)
                self.assertIsNone(text)
                self.assertTrue(error)
            router.assert_not_called()

    def test_http_options_are_forwarded_only_when_ai_is_requested(self):
        import app as application
        from pathlib import Path
        base = Path(__file__).resolve().parent

        def request_fields(enabled):
            return {"paper": (io.BytesIO((base/"examples/sample_paper.md").read_bytes()), "paper.md"),
                    "data": (io.BytesIO((base/"examples/student_scores.csv").read_bytes()), "data.csv"),
                    "use_llm": "1" if enabled else "0",
                    "ai_context_mode": "full", "ai_output_tokens": "16384",
                    "ai_research_tools": "1"}

        application.app.config["TESTING"] = True
        with patch.object(application, "audit_paper_with_llm", return_value=(
                "fixture-review", None, {"context": {"mode": "full"}})) as audit:
            client = application.app.test_client()
            r = client.post("/api/check_paper", data=request_fields(False))
            self.assertEqual(r.status_code, 200)
            audit.assert_not_called()
            r = client.post("/api/check_paper", data=request_fields(True))
            self.assertEqual(r.status_code, 200)
            self.assertEqual(audit.call_args.kwargs["context_mode"], "full")
            self.assertEqual(audit.call_args.kwargs["output_tokens"], "16384")
            self.assertTrue(audit.call_args.kwargs["research_tools"])
            self.assertEqual(r.get_json()["llm_context"], {"mode": "full"})

    def test_research_mode_bypasses_cached_review_and_attaches_tool_ledger(self):
        class Router:
            task_policy = None
            _last_usage = None
            _resolved = {}
            def _get_agent(self, state):
                return type("Agent", (), {"model_name": "fixture"})()
            def last_invocation(self):
                return {"actual_model": "custom_primary/fixture"}
        record = {"trace": [], "calls": [{"model": "custom_primary/fixture"}],
                  "sent_chars": 0, "usage": {"prompt_tokens": 20, "completion_tokens": 10}}
        with patch.object(agents, "get_router", return_value=Router()), \
             patch("research_review.run_research_review", return_value=("工具审阅", record)) as run, \
             patch.object(llm_audit.llm_cache, "get") as cache:
            text, error, meta = llm_audit.audit_paper_with_llm("测试论文", {}, "核对", research_tools=True)
        self.assertIsNone(error)
        self.assertIn("工具审阅", text)
        self.assertEqual(meta["research"], record)
        self.assertIsNone(meta["cost_estimate"])
        cache.assert_not_called()
        run.assert_called_once()

    def test_research_failure_keeps_known_usage_but_total_is_unknown(self):
        from research_review import ResearchReviewError
        class Router:
            task_policy = None
            def _get_agent(self, state):
                return type("Agent", (), {"model_name": "fixture"})()
        partial = {"calls": [{"usage": {"prompt_tokens": 100, "completion_tokens": 20}}],
                   "usage": None, "trace": []}
        with patch.object(agents, "get_router", return_value=Router()), \
             patch("research_review.run_research_review", side_effect=ResearchReviewError("timeout", partial)):
            text, error, meta = llm_audit.audit_paper_with_llm("论文", {}, "", research_tools=True)
        self.assertIsNone(text)
        self.assertTrue(error)
        self.assertIsNone(meta["input_tokens"])
        self.assertIsNone(meta["output_tokens"])
        self.assertEqual(meta["research"]["calls"][0]["usage"]["prompt_tokens"], 100)


if __name__ == "__main__":
    unittest.main()
