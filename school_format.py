"""本地学校格式核对。上传材料仅在请求内存中解析，不自动采信文件名或文本指令。"""
from __future__ import annotations

import hashlib
import io
import json
import math
import re
import zipfile
from pathlib import Path

from flask import Blueprint, jsonify, request
from docx import Document
from docx.oxml.ns import qn


school_format_bp = Blueprint("school_format", __name__)

MAX_REQUIREMENT = 8 * 1024 * 1024
MAX_REQUIREMENTS_TOTAL = 16 * 1024 * 1024
MAX_PAPER = 20 * 1024 * 1024
MAX_ZIP_UNPACKED = 80 * 1024 * 1024
MAX_SOURCE_TEXT = 200_000
FIELDS = {"east_asia_font", "latin_font", "font_size_pt", "line_spacing_multiplier",
          "first_line_indent_cm", "margin_top_cm", "margin_bottom_cm",
          "margin_left_cm", "margin_right_cm"}
FONT_FIELDS = {"east_asia_font", "latin_font"}
NUMERIC_RANGES = {"font_size_pt": (5, 72), "line_spacing_multiplier": (0.5, 4),
                  "first_line_indent_cm": (0, 5), "margin_top_cm": (0, 10),
                  "margin_bottom_cm": (0, 10), "margin_left_cm": (0, 10),
                  "margin_right_cm": (0, 10)}
FIELD_LABEL = {"east_asia_font": "中文字体", "latin_font": "西文字体", "font_size_pt": "字号",
               "line_spacing_multiplier": "行距倍数", "first_line_indent_cm": "首行缩进",
               "margin_top_cm": "上页边距", "margin_bottom_cm": "下页边距",
               "margin_left_cm": "左页边距", "margin_right_cm": "右页边距"}
BODY_START = re.compile(r"^(?:第[一二三四五六七八九十\d]+章(?:\s|$)|引言$|绪论$|前言$)")
NUMBERED_HEADING = re.compile(r"^[1-9]\d?(?:\.\d+)?\s+\S")
BODY_END = re.compile(r"^(?:参考文献|致谢|附录|References|Acknowledg(?:e)?ments)$", re.I)


def _read_limited(storage, limit):
    raw = storage.stream.read(limit + 1)
    if len(raw) > limit:
        raise ValueError(f"文件超过 {limit // (1024 * 1024)} MB 上限。")
    return raw


def _check_zip(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(x.file_size for x in entries) > MAX_ZIP_UNPACKED:
                raise ValueError("DOCX 解压内容过大，请转换或精简文件。")
            if any(x.flag_bits & 1 for x in entries):
                raise ValueError("加密 DOCX 暂不支持，请解除保护后重试。")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("文件不是有效 DOCX。")
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效 DOCX。") from exc


def _docx(raw):
    _check_zip(raw)
    try:
        return Document(io.BytesIO(raw))
    except Exception as exc:
        raise ValueError("DOCX 无法解析，请在 Word 中另存后重试。") from exc


def _source_text(raw, suffix):
    if suffix == ".docx":
        doc = _docx(raw)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError("加密 PDF 暂不支持，请解除保护后重试。")
            if len(reader.pages) > 80:
                raise ValueError("规范 PDF 超过 80 页，请拆分后重试。")
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF 无法解析，请转换成可提取文字的文件。") from exc
        if not text.strip():
            raise ValueError("PDF 未提取到文字，扫描件需要手动录入规则。")
    elif suffix in {".txt", ".md"}:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("文本编码无法识别。")
    else:
        raise ValueError("学校要求仅支持 DOCX、PDF、TXT、Markdown。")
    if len(text) > MAX_SOURCE_TEXT:
        raise ValueError("规范文字超过 20 万字，请拆分后重试。")
    return text.replace("\r\n", "\n"), suffix[1:]


def _load_sources(storages):
    out, total = [], 0
    for storage in storages:
        raw = _read_limited(storage, MAX_REQUIREMENT)
        total += len(raw)
        if total > MAX_REQUIREMENTS_TOTAL:
            raise ValueError("学校要求文件合计超过 16 MB。")
        name = Path(storage.filename or "").name
        text, fmt = _source_text(raw, Path(name).suffix.lower())
        out.append({"name": name, "sha256": hashlib.sha256(raw).hexdigest(),
                    "source_text": text, "format": fmt, "extraction_status": "text_only"})
    return out


def _safe_string(value, limit=200):
    if not isinstance(value, str) or len(value) > limit:
        raise ValueError("规则文字字段无效或过长。")
    return value.strip()


def _validate_profile(profile, sources):
    if not isinstance(profile, dict) or profile.get("confirmed") is not True:
        raise ValueError("请先确认本次规则及其适用范围。")
    scope = profile.get("scope")
    if not isinstance(scope, dict) or scope.get("body") is not True:
        raise ValueError("请确认规则适用于论文正文。")
    for key in ("school", "year", "degree"):
        _safe_string(scope.get(key, ""), 100)
    rules = profile.get("rules")
    if not isinstance(rules, list) or not rules or len(rules) > 40:
        raise ValueError("请提供 1 至 40 条本次确认的格式规则。")
    by_hash = {s["sha256"]: s for s in sources}
    cleaned = []
    for item in rules:
        if not isinstance(item, dict) or item.get("field") not in FIELDS:
            raise ValueError("规则字段不受支持。")
        if item.get("confirmed") is not True:
            raise ValueError("有规则尚未确认。")
        if type(item.get("active", True)) is not bool:
            raise ValueError("规则启用状态必须为布尔值。")
        field = item["field"]
        value = item.get("expected")
        if field in FONT_FIELDS:
            value = _safe_string(value, 100)
            if not value:
                raise ValueError("字体名称不能为空。")
        else:
            if isinstance(value, bool) or not isinstance(value, (float, int)):
                raise ValueError("数值规则需要明确数字。")
            value = float(value)
            low, high = NUMERIC_RANGES[field]
            if not math.isfinite(value) or not low <= value <= high:
                raise ValueError("规则数值超出支持范围。")
        source_type = item.get("source_type")
        if source_type not in {"school", "template", "custom", "default"}:
            raise ValueError("规则来源类型无效。")
        source_hash = item.get("source_hash")
        quote = _safe_string(item.get("quote", ""), 1000)
        if source_type in {"school", "template"}:
            source = by_hash.get(source_hash)
            if not source or not quote or quote not in source["source_text"]:
                raise ValueError("学校依据缺少本次上传文件中的原文，请重新选择并确认。")
        elif source_hash or quote:
            raise ValueError("自定义或通用参考规则不能标记为学校原文。")
        cleaned.append({"field": field, "expected": value, "source_type": source_type,
                        "source_hash": source_hash if source_type in {"school", "template"} else None,
                        "quote": quote, "active": item.get("active") is not False,
                        "reason": _safe_string(item.get("reason", ""), 300)})
    return cleaned


def _effective_rules(rules):
    """相互矛盾的正式学校规则暂停；自定义覆盖保留学校冲突标记。"""
    by_field = {}
    for rule in rules:
        if rule["active"]:
            by_field.setdefault(rule["field"], []).append(rule)
    selected, conflicts = {}, []
    priority = {"school": 3, "template": 2, "default": 1, "custom": 0}
    for field, entries in by_field.items():
        school = [r for r in entries if r["source_type"] == "school"]
        if len({str(r["expected"]) for r in school}) > 1:
            conflicts.append({"field": field, "status": "conflict", "location": "本次规则",
                              "expected": [r["expected"] for r in school], "actual": None,
                              "source": "学校要求之间冲突", "action": "核对适用范围和原文后再选择。",
                              "quote": [r["quote"] for r in school]})
            continue
        templates = [r for r in entries if r["source_type"] == "template"]
        if not school and len({str(r["expected"]) for r in templates}) > 1:
            conflicts.append({"field": field, "status": "conflict", "location": "本次规则",
                              "expected": [r["expected"] for r in templates], "actual": None,
                              "source": "官方模板说明之间冲突", "action": "核对适用范围和原文后再选择。",
                              "quote": [r["quote"] for r in templates]})
            continue
        custom = [r for r in entries if r["source_type"] == "custom"]
        if custom and school and any(r["expected"] != school[0]["expected"] for r in custom):
            if not custom[-1]["reason"]:
                raise ValueError("与学校要求冲突的自定义覆盖必须填写理由。")
            conflicts.append({"field": field, "status": "conflict", "location": "本次规则",
                              "expected": school[0]["expected"], "actual": custom[-1]["expected"],
                              "source": "自定义覆盖与学校要求冲突", "action": "查看学校原要求与覆盖理由；不能据此判定符合学校要求。",
                              "quote": school[0]["quote"], "reason": custom[-1]["reason"]})
            continue
        if len({str(r["expected"]) for r in custom}) > 1:
            conflicts.append({"field": field, "status": "conflict", "location": "本次规则",
                              "expected": [r["expected"] for r in custom], "actual": None,
                              "source": "自定义规则冲突", "action": "只保留本次适用的一条规则。"})
            continue
        # 无学校冲突时自定义可用于个人检查，但必须标记为个人选择。
        selected[field] = custom[-1] if custom else max(entries, key=lambda r: priority[r["source_type"]])
    return selected, conflicts


def _xml_font(element, key):
    if element is None:
        return None, False
    rpr = element.find(qn("w:rPr"))
    fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
    if fonts is None:
        return None, False
    direct = fonts.get(qn("w:" + key))
    themed = fonts.get(qn("w:" + key + "Theme"))
    if key == "ascii" and not direct:
        direct = fonts.get(qn("w:hAnsi"))
        themed = themed or fonts.get(qn("w:hAnsiTheme"))
    return direct, bool(themed)


def _style_chain(run, paragraph, document):
    visited = set()
    if run.style is not None:
        style = run.style
        while style is not None and style.style_id not in visited:
            visited.add(style.style_id)
            yield style._element
            style = style.base_style
    style = paragraph.style
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        yield style._element
        style = style.base_style
    defaults = document.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rdef = defaults.find(qn("w:rPrDefault"))
        if rdef is not None:
            yield rdef


def _effective_font(run, paragraph, document, key):
    for element in (run._element, *_style_chain(run, paragraph, document)):
        value, themed = _xml_font(element, key)
        if value:
            return value
        if themed:
            return None
    return None


def _effective_size(run, paragraph, document):
    for element in (run._element, *_style_chain(run, paragraph, document)):
        rpr = element.find(qn("w:rPr"))
        if rpr is None:
            continue
        size = rpr.find(qn("w:sz"))
        if size is not None and size.get(qn("w:val")):
            try:
                return int(size.get(qn("w:val"))) / 2
            except ValueError:
                return None
    return None


def _paragraph_format(paragraph, field):
    styles = [paragraph.style]
    seen = set()
    while styles[-1] is not None:
        style = styles[-1]
        if style.style_id in seen:
            break
        seen.add(style.style_id)
        next_style = style.base_style
        if next_style is None or next_style.style_id in seen:
            break
        styles.append(next_style)
    for fmt in (paragraph.paragraph_format, *(s.paragraph_format for s in styles if s is not None)):
        if field == "first_line_indent":
            ppr = fmt._element.find(qn("w:pPr"))
            ind = ppr.find(qn("w:ind")) if ppr is not None else None
            if ind is not None and ind.get(qn("w:firstLineChars")) is not None:
                return "character_units"
        value = getattr(fmt, field)
        if value is not None:
            return value
    return None


def _value(document, paragraph, run, field, section=None):
    if field == "east_asia_font":
        return _effective_font(run, paragraph, document, "eastAsia")
    if field == "latin_font":
        return _effective_font(run, paragraph, document, "ascii")
    if field == "font_size_pt":
        return _effective_size(run, paragraph, document)
    if field == "line_spacing_multiplier":
        value = _paragraph_format(paragraph, "line_spacing")
        return float(value) if isinstance(value, float) else None
    if field == "first_line_indent_cm":
        value = _paragraph_format(paragraph, "first_line_indent")
        return round(value.cm, 4) if value is not None and value != "character_units" else None
    if field.startswith("margin_") and section is not None:
        edge = field.removeprefix("margin_").removesuffix("_cm")
        value = getattr(section, edge + "_margin")
        return round(value.cm, 4) if value is not None else None
    return None


def _matches(actual, expected, field):
    if field in FONT_FIELDS:
        return actual.strip().casefold() == expected.strip().casefold()
    tolerance = 0.01 if field == "font_size_pt" else 0.005
    return abs(actual - expected) <= tolerance


def _finding(status, location, field, rule, actual, action):
    return {"status": status, "location": location, "field": field,
            "label": FIELD_LABEL[field], "expected": rule["expected"], "actual": actual,
            "source": rule["source_type"], "source_hash": rule["source_hash"],
            "quote": rule["quote"], "reason": rule["reason"], "action": action}


def _body_paragraphs(document):
    started = False
    for index, para in enumerate(document.paragraphs, 1):
        content = para.text.strip()
        if not content:
            continue
        if BODY_END.match(content):
            break
        if not started:
            style_name = para.style.name if para.style is not None else ""
            styled_heading = "heading" in style_name.lower() or "标题" in style_name
            if len(content) < 100 and (BODY_START.match(content) or
                                       (styled_heading and NUMBERED_HEADING.match(content))):
                started = True
            continue
        style_name = para.style.name if para.style is not None else ""
        if "heading" in style_name.lower() or "标题" in style_name:
            continue
        if BODY_START.match(content) and len(content) < 100:
            continue
        yield index, para


def _check_docx(raw, rules):
    document = _docx(raw)
    findings = []
    body = list(_body_paragraphs(document))
    max_findings = 5000
    truncated = False
    body_checked = 0
    content_fields = [f for f in rules if not f.startswith("margin_")]
    if not body and content_fields:
        for field in content_fields:
            findings.append(_finding("undetermined", "论文正文", field, rules[field], None,
                                     "无法可靠定位正文起点，请人工核对或添加章节标题。"))
    for index, paragraph in body[:2000]:
        body_checked += 1
        for field in content_fields:
            rule = rules[field]
            if field in FONT_FIELDS or field == "font_size_pt":
                runs = [(n, r) for n, r in enumerate(paragraph.runs, 1) if r.text.strip()]
            else:
                runs = [(None, None)]
            evaluated = False
            for run_index, run in runs:
                if field == "east_asia_font" and not re.search(r"[\u3400-\u9fff]", run.text):
                    continue
                if field == "latin_font" and not re.search(r"[A-Za-z]", run.text):
                    continue
                if run is not None and (run.font.superscript or run.font.subscript):
                    continue
                evaluated = True
                location = f"正文第{index}段" + (f"第{run_index}片段" if run_index else "")
                actual = _value(document, paragraph, run, field)
                status = "undetermined" if actual is None else ("compliant" if _matches(actual, rule["expected"], field) else "noncompliant")
                action = "继承属性或主题字体无法确定，请在 Word 中核对。" if actual is None else (
                    "无需修改。" if status == "compliant" else f"请在 Word 中核对并设置{FIELD_LABEL[field]}。")
                findings.append(_finding(status, location, field, rule, actual, action))
                if len(findings) >= max_findings:
                    truncated = True
                    break
            if not evaluated:
                findings.append(_finding("not_applicable", f"正文第{index}段", field, rule, None,
                                         "该段没有适用文字片段，或片段属于已排除的上下标。"))
            if truncated:
                break
        if truncated:
            break
    for i, section in enumerate(document.sections, 1):
        for field, rule in rules.items():
            if not field.startswith("margin_"):
                continue
            actual = _value(document, None, None, field, section)
            status = "undetermined" if actual is None else ("compliant" if _matches(actual, rule["expected"], field) else "noncompliant")
            findings.append(_finding(status, f"第{i}节页面", field, rule, actual,
                                     "无法读取分节页边距，请人工核对。" if actual is None else (
                                         "无需修改。" if status == "compliant" else "请在 Word 的页面设置中核对该节页边距。")))
    coverage = {"body_paragraphs_checked": body_checked, "body_paragraphs_total": len(body),
                "sections_checked": len(document.sections),
                "excluded": ["封面、摘要、标题、参考文献与致谢的格式", "表格、页眉页脚、脚注、文本框与公式",
                             "自动分页及实际页码", "无法解析的主题字体与继承属性"]}
    if len(body) > 2000:
        coverage["excluded"].append("正文第 2001 段及以后")
    if truncated:
        coverage["excluded"].append("达到 5000 条结果上限后的正文属性")
    return findings, coverage


def _markdown(profile_version, paper_hash, findings, coverage):
    lines = ["# 学校规范与 DOCX 格式核对", "", f"规则档案版本：{profile_version}",
             f"论文 SHA256：{paper_hash}", "", "结果仅覆盖列出的规则和区域；未判定不等于通过。", ""]
    for finding in findings:
        label = FIELD_LABEL.get(finding.get("field"), finding.get("field", "规则"))
        lines.append(f"## {finding['status']} · {label} · {finding['location']}")
        lines.append(f"- 要求：{finding['expected']}；实际：{finding['actual']}")
        lines.append(f"- 来源：{finding['source']}；依据：{str(finding.get('quote') or '')[:200]}")
        if finding.get('reason'):
            lines.append(f"- 用户覆盖理由：{finding['reason']}（不据此表示符合学校要求）")
        lines.append(f"- 建议：{finding['action']}")
        lines.append("")
    lines.append("未覆盖：" + "；".join(coverage["excluded"]))
    return "\n".join(lines)


@school_format_bp.route("/api/school_rules/preview", methods=["POST"])
def preview_school_rules():
    """展示上传文件中的文字供用户手动确认；不自动生成或启用规则。"""
    try:
        sources = _load_sources(request.files.getlist("requirements"))
        if not sources:
            raise ValueError("请单独上传学校规范文件。")
        return jsonify({"ok": True, "sources": sources,
                        "notes": ["文件名不证明官方身份或适用性。", "原文仅供核对；确认规则后才参与检查。"]})
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


@school_format_bp.route("/api/format_check", methods=["POST"])
def format_check():
    """独立论文与学校要求上传；所有材料在本次请求内读取和核对。"""
    try:
        paper = request.files.get("paper")
        if paper is None:
            raise ValueError("请上传论文文件。")
        name = Path(paper.filename or "").name
        suffix = Path(name).suffix.lower()
        if suffix not in {".docx", ".pdf"}:
            raise ValueError("格式属性检查仅支持 DOCX；PDF 返回未判定。")
        raw = _read_limited(paper, MAX_PAPER)
        sources = _load_sources(request.files.getlist("requirements"))
        try:
            profile = json.loads(request.form.get("profile", ""))
        except (ValueError, TypeError) as exc:
            raise ValueError("规则档案 JSON 无效。") from exc
        rules = _validate_profile(profile, sources)
        effective, conflicts = _effective_rules(rules)
        if suffix == ".docx":
            findings, coverage = _check_docx(raw, effective)
        else:
            try:
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(raw))
                if reader.is_encrypted:
                    raise ValueError("加密 PDF 暂不支持，请解除保护后重试。")
                if len(reader.pages) > 1000:
                    raise ValueError("论文 PDF 页数过多，请拆分后重试。")
            except ValueError:
                raise
            except Exception as exc:
                raise ValueError("PDF 无法解析，请重新导出后重试。") from exc
            findings = [_finding("undetermined", "PDF 全文", f, rule, None,
                                 "PDF 不保留可复核的 Word 样式；请上传 DOCX 或人工核对。")
                        for f, rule in effective.items()]
            coverage = {"body_paragraphs_checked": 0, "body_paragraphs_total": 0,
                        "sections_checked": 0, "excluded": ["PDF 字体、字号、段落及页边距格式未检查"]}
        for rule in rules:
            if rule["active"] is False:
                findings.append(_finding("not_applicable", "本次关闭的规则", rule["field"], rule, None,
                                         "启用后重新检查。"))
        findings = conflicts + findings
        # 源文件哈希和确认规则决定版本；旧报告不会因下次编辑而变化。
        version_payload = {"scope": profile["scope"], "rules": rules}
        profile_version = hashlib.sha256(json.dumps(version_payload, sort_keys=True,
                                                    ensure_ascii=False).encode()).hexdigest()[:16]
        paper_hash = hashlib.sha256(raw).hexdigest()
        counts = {status: sum(f["status"] == status for f in findings)
                  for status in ("compliant", "noncompliant", "undetermined", "not_applicable", "conflict")}
        provenance = {"scope": profile["scope"], "sources": [
            {"name": s["name"], "sha256": s["sha256"]} for s in sources], "rules": rules}
        markdown = _markdown(profile_version, paper_hash, findings, coverage)
        markdown += "\n\n## 本次规则档案与适用范围\n\n" + json.dumps(provenance, ensure_ascii=False, indent=2)
        return jsonify({"ok": True, "paper_hash": paper_hash, "profile_version": profile_version,
                        "findings": findings, "counts": counts, "coverage": coverage,
                        "provenance": provenance, "markdown": markdown,
                        "notes": ["仅核对用户确认的格式规则；自定义规则不代表学校要求。",
                                  "未判定及未覆盖范围不等于符合。"]})
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
