"""v0.9 · Word 导出端到端测试。"""
import base64
import io as _io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from app import app, run_independent_t, run_anova, run_paired_t, _build_chart_png  # noqa: E402

import pandas as pd  # noqa: E402


def main() -> None:
    df = pd.read_csv(Path(__file__).parent / "examples/student_scores.csv")
    client = app.test_client()

    print("=" * 70)
    print("7 个方法导出 docx")
    print("=" * 70)

    methods = [
        ("independent_t", ("gender", "score", None)),
        ("anova", ("study_intensity", "score", None)),
        ("correlation", (None, "study_hours", "score")),
        ("chi_square", ("gender", "pass", None)),
        ("paired_t", (None, "anxiety_pre", "anxiety_post")),
        ("mann_whitney", ("gender", "reaction_time_ms", None)),
        ("wilcoxon", (None, "anxiety_pre", "anxiety_post")),
    ]

    from app import (run_correlation, run_chi_square,  # noqa: E402
                     run_mann_whitney, run_wilcoxon)
    run_fns = {
        "independent_t": run_independent_t, "anova": run_anova,
        "correlation": run_correlation, "chi_square": run_chi_square,
        "paired_t": run_paired_t, "mann_whitney": run_mann_whitney,
        "wilcoxon": run_wilcoxon,
    }

    for method, chart_args in methods:
        g, v1, v2 = chart_args
        fn = run_fns[method]
        if method in ("correlation", "paired_t", "wilcoxon"):
            result = fn(df, v1, v2)
        else:
            result = fn(df, g, v1)

        png, _ = _build_chart_png(df, method, g, v1, v2)
        b64 = base64.b64encode(png).decode() if png else None

        body = {"markdown": result["markdown"], "method": method,
                "meta": {"数据文件": "student_scores.csv", "n": 30}}
        if b64:
            body["chart_base64"] = b64
        resp = client.post("/api/export", json=body)
        assert resp.status_code == 200, f"{method}: HTTP {resp.status_code}"
        assert len(resp.data) > 5000, f"{method}: docx 太小 {len(resp.data)}"

        # 验证 docx 能重新打开
        from docx import Document
        doc = Document(_io.BytesIO(resp.data))
        assert len(doc.paragraphs) > 5, f"{method}: 段落太少"
        assert len(doc.tables) >= 1, f"{method}: 缺表格"

        print(f"  ✓ {method:<14} {len(resp.data):>7} bytes, "
              f"段落={len(doc.paragraphs)}, 表格={len(doc.tables)}")

    print()
    print("=" * 70)
    print("错误路径")
    print("=" * 70)
    resp = client.post("/api/export", json={"method": "independent_t"})
    assert resp.status_code == 400
    print(f"  ✓ 缺 markdown → 400 {resp.get_json()['error']}")

    resp = client.post("/api/export", json={"markdown": "## x", "method": "t",
                                             "chart_base64": "@@@bad@@@"})
    assert resp.status_code == 200  # 坏图不阻断
    print(f"  ✓ 坏 base64 → 200（图跳过，{len(resp.data)} bytes）")

    resp = client.post("/api/export", json={
        "markdown": "## t\n\n| a | b |\n| --- | --- |\n| 1 | 2 |\n\n- item\n\npara",
        "method": "t"})
    assert resp.status_code == 200
    from docx import Document
    doc = Document(_io.BytesIO(resp.data))
    assert len(doc.tables) == 1
    print(f"  ✓ 简单 markdown（标题+表+列表+段）→ 表格={len(doc.tables)}")

    # 中文文件名
    resp = client.post("/api/export", json={"markdown": "## 报告", "method": "independent_t"})
    cd = resp.headers.get("Content-Disposition", "")
    assert "attachment" in cd
    print(f"  ✓ 附件下载头存在")

    print()
    print("=" * 70)
    print("✅ Word 导出端到端测试全部通过")
    print("=" * 70)


if __name__ == "__main__":
    main()


def test_audit_export() -> None:
    """v0.9.5：论文排查报告导出（title 参数 + 无 method 路径 + 文件名）。"""
    import pandas as pd
    from pathlib import Path as P
    client = app.test_client()

    # 走一遍真实论文排查接口（sample_paper.md + 示例数据）
    paper = P(__file__).parent / "examples" / "sample_paper.md"
    data = P(__file__).parent / "examples" / "student_scores.csv"
    r = client.post("/api/check_paper",
                    data={"paper": (open(paper, "rb"), "p.md"),
                          "data": (open(data, "rb"), "d.csv")},
                    content_type="multipart/form-data")
    d = r.get_json()
    assert d.get("ok"), f"check_paper 失败: {d.get('error')}"

    md = d["audit"]["markdown"]
    assert md, "audit markdown 为空"

    # 无 method、带 title 的导出
    resp = client.post("/api/export", json={
        "markdown": md,
        "title": "论文排查报告",
        "meta": {"论文文件": "p.md", "数据文件": "d.csv", "样本量": 30},
    })
    assert resp.status_code == 200, f"HTTP {resp.status_code}"
    assert len(resp.data) > 5000, f"docx 太小 {len(resp.data)}"
    # 附件文件名应含"论文排查报告"（HTTP header 中文为 UTF-8 百分号编码）
    cd = resp.headers.get("Content-Disposition", "")
    import urllib.parse
    cd_decoded = urllib.parse.unquote(cd)
    assert "论文排查报告" in cd_decoded, f"文件名不含论文排查报告: {cd_decoded}"

    from docx import Document
    doc = Document(_io.BytesIO(resp.data))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert any("论文排查报告" in t for t in texts), "主标题应为论文排查报告"
    assert any("智论助手" in t for t in texts), "标题应含智论助手"

    # 旧路径回归：无 title 无 method（兜底"报告"，不崩）
    resp2 = client.post("/api/export", json={"markdown": "## 测试\n\n内容"})
    assert resp2.status_code == 200
    print("  [PASS] 论文排查报告导出（title + 文件名 + 兜底路径）")


if __name__ == "__main__":
    test_audit_export()
    print("export_test: 全部通过")
