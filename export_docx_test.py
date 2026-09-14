"""v2.18 · `export_docx` 模块级回归测试。

为什么单独开一个文件：原有的 `export_test.py` 只做**接口冒烟**（7 个方法能导出、
docx 能重开），一个断言都不计数，也从不碰边界。结果 `export_docx.py` 这个 v0.9
"结果可导出"的核心交付物，藏着四个缺陷一直没人发现：

  1. `#### 四级标题` 不被识别，原样输出成 "#### xxx"；
  2. 整行都是 `-` 的**数据行**被当成 GFM 分隔线吞掉 —— 静默丢一行数据；
  3. 数据行比表头多列时静默截断 —— 多出来的数字直接消失；
  4. `markdown=None` 抛 `AttributeError`。

对论文工具来说，**静默丢数字**比报错糟得多：用户拿到一份"看起来正常"的 docx，
交上去才发现表里少了一行。所以本测试的主线不是"能不能导出"，而是
**"进去的每个数字，出来还在不在"**。

跑法：python export_docx_test.py
退出码：0 全通过 / 1 有真失败 / 2 环境未就绪（缺 python-docx）
"""
from __future__ import annotations

import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    from docx import Document
except ImportError:  # 环境未就绪，不是 bug
    print("export_docx_test: 未安装 python-docx，跳过（rc=2）")
    sys.exit(2)

import export_docx  # noqa: E402

PASS = 0
FAIL = 0
SKIP = 0
_LINES: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        _LINES.append(f"  [PASS] {name}")
    else:
        FAIL += 1
        _LINES.append(f"  [FAIL] {name}" + (f"  {detail}" if detail else ""))


def section(title: str) -> None:
    _LINES.append("")
    _LINES.append("=" * 72)
    _LINES.append(title)
    _LINES.append("=" * 72)


def skip(name: str, why: str) -> None:
    global SKIP
    SKIP += 1
    _LINES.append(f"  [SKIP] {name}  {why}")


# ---------------------------------------------------------------------------
# 读取辅助
# ---------------------------------------------------------------------------
def open_doc(blob: bytes) -> Document:
    return Document(io.BytesIO(blob))


def all_text(doc: Document) -> str:
    """段落 + 表格里所有文字，拼成一个大串——用来查"某个数字还在不在"。"""
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def table_grid(doc: Document) -> list[list[list[str]]]:
    return [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]


def headings(doc: Document) -> list[tuple[str, str]]:
    """[(样式名, 文本)]，只取标题类段落。"""
    out = []
    for p in doc.paragraphs:
        if p.style and p.style.name.startswith(("Heading", "Title")):
            out.append((p.style.name, p.text))
    return out


def render(md, **kw) -> Document:
    return open_doc(export_docx.markdown_to_docx(md, **kw))


# ---------------------------------------------------------------------------
# 1. 标题层级
# ---------------------------------------------------------------------------
section("1. 标题层级（# ~ ######，v2.18 修复①：四级及以下不再原样输出）")

md_head = "\n\n".join([
    "# 一级标题", "## 二级标题", "### 三级标题",
    "#### 四级标题", "##### 五级标题", "###### 六级标题",
    "####只有井号没有空格不算标题",
])
doc = render(md_head)
hs = headings(doc)
texts = [t for _, t in hs]

for lv, txt in [(1, "一级标题"), (2, "二级标题"), (3, "三级标题"),
                (4, "四级标题"), (5, "五级标题"), (6, "六级标题")]:
    check(f"{'#' * lv} 标题文字进文档", txt in texts, f"实际标题={texts}")

body = all_text(doc)
for lv in range(1, 7):
    check(f"{'#' * lv} 不残留井号", f"{'#' * lv} " not in body, body[:200])

styles = {t: s for s, t in hs}  # {标题文字: 样式名}
check("#### 识别为 Heading 3", styles.get("四级标题") == "Heading 3", str(styles))
check("### 识别为 Heading 2（历史行为保留）", styles.get("三级标题") == "Heading 2", str(styles))
check("## 识别为 Heading 1（历史行为保留）", styles.get("二级标题") == "Heading 1", str(styles))
check("##### 识别为 Heading 4", styles.get("五级标题") == "Heading 4", str(styles))
check("###### 识别为 Heading 4（封顶）", styles.get("六级标题") == "Heading 4", str(styles))
check("井号后无空格不当标题（当普通段落）",
      "####只有井号没有空格不算标题" in body and
      "####只有井号没有空格不算标题" not in [t for _, t in hs])

# ---------------------------------------------------------------------------
# 2. 表格：结构与数据守恒
# ---------------------------------------------------------------------------
section("2. 表格结构")

md_t = "| 变量 | 均值 | 标准差 |\n| --- | --- | --- |\n| 年龄 | 20.5 | 3.1 |\n| 得分 | 78.2 | 9.4 |"
grid = table_grid(render(md_t))
check("生成 1 张表", len(grid) == 1, str(grid))
check("表头 3 列", grid[0][0] == ["变量", "均值", "标准差"], str(grid[0][0]))
check("数据 2 行（不含表头）", len(grid[0]) == 3, str(grid))
check("第 1 行内容正确", grid[0][1] == ["年龄", "20.5", "3.1"], str(grid[0][1]))
check("第 2 行内容正确", grid[0][2] == ["得分", "78.2", "9.4"], str(grid[0][2]))
check("分隔行不变成数据行", all("---" not in c for row in grid[0] for c in row), str(grid))

section("3. v2.18 修复②：整行是 `-` 的数据行不被吞")

md_dash = ("| 变量 | 值 |\n| --- | --- |\n| A | 1 |\n| - | - |\n| C | 3 |")
grid = table_grid(render(md_dash))
check("3 行数据都在（含全 `-` 行）", len(grid[0]) == 4, str(grid))
check("全 `-` 行原样保留", ["-", "-"] in grid[0], str(grid))
check("全 `-` 行没被当分隔行删掉", any(r == ["-", "-"] for r in grid[0]), str(grid))

grid = table_grid(render("| a | b |\n| --- | --- |"))
check("表头 + 分隔行 → 只剩表头一行", len(grid[0]) == 1, str(grid))
grid3 = table_grid(render("| a | b |\n| --- | --- |\n| x | y |\n| --- | --- |"))
check("非首位的 `---` 行当数据保留", len(grid3[0]) == 3 and grid3[0][2] == ["---", "---"],
      str(grid3))

section("4. v2.18 修复③：数据行比表头多列 → 扩列，不静默截断")

md_wide = "| a | b |\n| --- | --- |\n| 1 | 2 | 3 |\n| 4 | 5 |"
grid = table_grid(render(md_wide))
check("列数按最宽行扩到 3", len(grid[0][0]) == 3, str(grid[0][0]))
check("多出来的 3 没丢", grid[0][1] == ["1", "2", "3"], str(grid[0][1]))
check("短行补空不报错", grid[0][2][:2] == ["4", "5"], str(grid[0][2]))
check("表头被补空而非丢列", grid[0][0][:2] == ["a", "b"], str(grid[0][0]))

section("5. 空单元格")

grid = table_grid(render("| 项目 | 前测 | 后测 |\n| --- | --- | --- |\n| A |  | 5 |"))
check("中间空格子保留（3 列）", grid[0][1] == ["A", "", "5"], str(grid[0][1]))
grid = table_grid(render("|  | x |  |\n| --- | --- | --- |\n| 1 | 2 | 3 |"))
check("首尾空格子不被 strip('|') 吃掉", grid[0][0] == ["", "x", ""], str(grid[0][0]))

# ---------------------------------------------------------------------------
# 6. 行内标记与列表
# ---------------------------------------------------------------------------
section("6. 粗体 / 列表 / 段落")

doc = render("**均值**：3.5\n\n- 第一项\n* 第二项\n\n普通段落")
body = all_text(doc)
check("粗体内容在", "均值" in body and "3.5" in body, body[:200])
check("粗体标记不残留 **", "**" not in body, body[:200])
bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
check("存在加粗 run", "均值" in bold_runs, str(bold_runs))
check("列表项文字在", "第一项" in body and "第二项" in body)
bullet = [p for p in doc.paragraphs if p.style and p.style.name == "List Bullet"]
check("两项都用了 List Bullet 样式", len(bullet) == 2, str([p.text for p in bullet]))
check("普通段落在", "普通段落" in body)

# ---------------------------------------------------------------------------
# 7. 数字守恒（主线）
# ---------------------------------------------------------------------------
section("7. 数字守恒：进去的每个数字，出来还在")

NUMS = ["20.500", "3.100", "78.200", "9.400", "-2.431", "0.0153", "0.9876",
        "12", "-0.75", "1.0e-4", "99", "0.001"]
md_nums = (
    "## 结果\n\n"
    "| 指标 | 数值 |\n| --- | --- |\n"
    + "".join(f"| 指标{i} | {v} |\n" for i, v in enumerate(NUMS))
    + "\n结论：**t** = -2.431，p = 0.0153。\n"
)
body = all_text(render(md_nums))
for v in NUMS:
    check(f"数字 {v} 出现在 docx 里", v in body, body[:300])

# ---------------------------------------------------------------------------
# 8. 图表嵌入
# ---------------------------------------------------------------------------
section("8. 图表嵌入")

PNG_1X1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000a49444154789c6300010000050001"
    "0d0a2db40000000049454e44ae426082"
)
doc = render("## 图\n", chart_png=PNG_1X1)
body = all_text(doc)
check("合法 PNG 不报失败", "（图表嵌入失败）" not in body, body[:200])
check("有图注", "（数据可视化图表）" in body)
has_pic = any(p._element.findall(
    ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    for p in doc.paragraphs)
check("文档里真的有图片节点", has_pic)

for bad, label in [(b"not a png", "非法字节"), ("字符串", "非 bytes"), (b"", "空 bytes")]:
    d = render("## 图\n", chart_png=bad)
    check(f"{label}的图不抛异常", isinstance(all_text(d), str))
check("非法 PNG 给出降级提示", "（图表嵌入失败）" in all_text(render("## 图\n", chart_png=b"x")))
check("无图时不出现图注", "（数据可视化图表）" not in all_text(render("## 图\n")))

# ---------------------------------------------------------------------------
# 9. 脏输入（v2.18 修复④ + 周边宽容度）
# ---------------------------------------------------------------------------
section("9. 脏输入：纯函数应当宽容，不抛异常")

DIRTY = [
    ("markdown=None", None),
    ("markdown=''", ""),
    ("markdown=123", 123),
    ("markdown=列表", ["a", "b"]),
    ("markdown=bytes", "二进制".encode("utf-8")),
    ("markdown=只有换行", "\n\n\n"),
    ("markdown=只有分隔行", "| --- | --- |"),
    ("markdown=未闭合表格", "| a | b |\n| 1 | 2 |"),
]
for name, val in DIRTY:
    try:
        blob = export_docx.markdown_to_docx(val)
        ok = isinstance(blob, bytes) and blob[:2] == b"PK" and len(blob) > 1000
        check(f"{name} → 合法 docx", ok, f"len={len(blob)} head={blob[:4]!r}")
    except Exception as exc:  # noqa: BLE001
        check(f"{name} → 合法 docx", False, f"{type(exc).__name__}: {exc}")

check("markdown=123 内容转成了 '123'", "123" in all_text(render(123)))
check("markdown=bytes 能解码", "二进制" in all_text(render("二进制".encode("utf-8"))))

for name, kw in [
    ("title=None", {"title": None}),
    ("title=123", {"title": 123}),
    ("meta=列表", {"meta": ["x"]}),
    ("meta=None", {"meta": None}),
    ("method_label=None", {"method_label": None}),
    ("method_label=123", {"method_label": 123}),
]:
    try:
        blob = export_docx.markdown_to_docx("## x", **kw)
        check(f"{name} 不抛异常", isinstance(blob, bytes) and len(blob) > 1000)
    except Exception as exc:  # noqa: BLE001
        check(f"{name} 不抛异常", False, f"{type(exc).__name__}: {exc}")

# ---------------------------------------------------------------------------
# 10. 封面 / 元信息 / 免责声明
# ---------------------------------------------------------------------------
section("10. 封面与固定文案")

doc = render("## x", method_label="独立样本 T 检验",
             meta={"数据文件": "a.csv", "样本量": 30, "空值": None},
             title="自定义标题")
body = all_text(doc)
check("主标题在", "自定义标题" in body)
check("方法副标题在", "方法：独立样本 T 检验" in body)
check("元信息在", "a.csv" in body and "30" in body)
check("元信息跳过空值", "空值" not in body)
check("免责声明在", "仅作数据辅助分析参考" in body)

doc2 = render("## x")
check("默认主标题在", "智论助手 · 统计分析报告" in all_text(doc2))
check("无 meta 时正文仍完整（免责声明在）", "仅作数据辅助分析参考" in all_text(doc2))
check("无 method_label 时不写方法行", "方法：" not in all_text(doc2))

# ---------------------------------------------------------------------------
# 11. 集成：真实报告（数字守恒 + 表格在）
# ---------------------------------------------------------------------------
section("11. 集成：真实统计报告导出")

try:
    import pandas as pd

    import app as app_mod
except Exception as exc:  # noqa: BLE001
    skip("真实报告导出", f"app/pandas 不可导入：{exc}")
else:
    csv_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "examples", "student_scores.csv")
    if not os.path.exists(csv_path):
        skip("真实报告导出", f"缺示例数据 {csv_path}")
    else:
        df = pd.read_csv(csv_path)
        cases = [
            ("independent_t", lambda: app_mod.run_independent_t(df, "gender", "score")),
            ("anova", lambda: app_mod.run_anova(df, "study_intensity", "score")),
            ("correlation", lambda: app_mod.run_correlation(df, "study_hours", "score")),
            ("paired_t", lambda: app_mod.run_paired_t(df, "anxiety_pre", "anxiety_post")),
        ]
        for name, fn in cases:
            try:
                res = fn()
            except Exception as exc:  # noqa: BLE001
                skip(f"{name} 导出", f"run 失败：{exc}")
                continue
            md = res.get("markdown") or ""
            if not md:
                skip(f"{name} 导出", "markdown 为空")
                continue
            # 报告里出现的数字，导出后必须还在
            import re as _re
            nums = [n for n in _re.findall(r"-?\d+\.\d{2,}", md)]
            nums = list(dict.fromkeys(nums))[:8]
            d = render(md, method_label=name)
            txt = all_text(d)
            lost = [n for n in nums if n not in txt]
            check(f"{name}：{len(nums)} 个数字零丢失", not lost, f"丢了 {lost}")
            check(f"{name}：至少 1 张表", len(d.tables) >= 1, str(len(d.tables)))
            check(f"{name}：无 ** 残留", "**" not in txt)

# ---------------------------------------------------------------------------
out = "\n".join(_LINES)
sys.stdout.buffer.write(out.encode("utf-8"))
sys.stdout.buffer.write(
    f"\n\nexport_docx_test: {PASS} 通过 / {FAIL} 失败 / {SKIP} 跳过\n".encode("utf-8"))
sys.exit(1 if FAIL else 0)
